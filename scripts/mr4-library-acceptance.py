"""Saved documents (MR4), live: a real model answers from a person's own files.

The unit checks drive Workbench against a fake gateway. What only a live run shows is
the whole path a document takes: a real Workbench process reading each file in a child
process of its own, indexing it, being stopped in the middle and starting again; a real
gateway, a real inference-driver and a real `llama-server` answering from a real small
model that has never seen the code words in the files. Two recording proxies keep what
each hop was sent: what Workbench sent the gateway, and what reached the engine. The
people are real sign-ins against a stand-in for Eugene's sign-in (OpenID Connect, as the
unit checks use).

  0. isolated: every ambient EUGENE_PLEXUS_* variable dropped, owned ports free
  1. the stack: engine (CPU, Qwen3 0.6B, a window of 8,192), a driver that says it is
     local, a gateway, Workbench; two people and the owner sign in
  2. three files into My documents: each file's own state while they are read, then
     their counts and the totals
  3. a question whose answer is a made-up code word on page 4 of one file: the engine is
     sent the passage (latest message only, before the question), the real model gives
     the word, and the answer's Sources name the file and the page; Test a question and
     Show indexed text say the same, and ask no model
  4. edit (add the file again) and delete: the old words are gone from answers at once
  5. Workbench is killed in the middle of reading twelve files and starts again: each is
     read once, none twice, none left saying "reading"
  6. a PDF with no text, and a locked one, are skipped with their reasons
  7. nothing is nothing matched: a question the files do not answer says so, and a search
     that cannot run says that instead, in its own words, and Rebuild brings it back
  8. documents only: Workbench answers in its own words and asks no model
  9. a file attached to a chat and searched instead: the model answers from it, and it
     goes with the chat
  10. two people on every endpoint, and the owner
  11. teardown by pid; no owned port still listening

**Safe beside a live install**: every ambient EUGENE_PLEXUS_* variable is dropped, ports
are 8480-8487, teardown is by pid. Never `pkill -f eugene_plexus_`: this box is a worker
node.

Run with the Workbench interpreter (it has httpx, fastapi, uvicorn, joserfc, pypdf):
  workbench/.venv/Scripts/python.exe workbench-mr/scripts/mr4-library-acceptance.py
"""

from __future__ import annotations

import base64
import contextlib
import hashlib
import http.client
import json
import os
import re
import secrets
import socket
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import traceback
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any, cast
from urllib.parse import urlencode, urlsplit

import httpx
import uvicorn
from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse, RedirectResponse
from joserfc import jwt
from joserfc.jwk import RSAKey

ROOT = Path(os.environ.get("EP_ROOT", "D:/py/eugene-plexus"))
HERE = Path(__file__).resolve().parent.parent
#: The Workbench tree under test: the one this script sits in, else the main one.
WB = Path(
    os.environ.get(
        "EP_WORKBENCH",
        str(HERE if (HERE / "src" / "eugene_plexus_workbench").is_dir() else ROOT / "workbench"),
    )
)
WB_PY = Path(
    os.environ.get("EP_WORKBENCH_PY", str(ROOT / "workbench" / ".venv" / "Scripts" / "python.exe"))
)
GW_PY = ROOT / "gateway" / ".venv" / "Scripts" / "python.exe"
DRV_PY = ROOT / "inference-driver" / ".venv" / "Scripts" / "python.exe"
TEMP = Path(os.environ.get("TEMP", tempfile.gettempdir()))
LLAMA = Path(
    os.environ.get(
        "EP_LLAMA_SERVER",
        str(
            Path.home() / ".eugene-plexus" / "engines" / "llama_cpp" / "b10948" / "llama-server.exe"
        ),
    )
)
MODEL_FILE = Path(
    os.environ.get("EP_MODEL_FILE", str(TEMP / "ep-s9-model" / "Qwen3-0.6B-Q4_K_M.gguf"))
)

sys.path.insert(0, str(HERE))
from tests import documents as docs  # noqa: E402  (the fixtures the unit checks use)

(
    GW_PORT,
    AGENT_PORT,
    DRV_LOCAL,
    ENGINE,
    ENGINE_PROXY,
    WB_PORT,
    EUGENE_PORT,
    GW_PROXY,
) = range(8480, 8488)
OWNED = [GW_PORT, AGENT_PORT, DRV_LOCAL, ENGINE, ENGINE_PROXY, WB_PORT, EUGENE_PORT, GW_PROXY]
#: Ports a thread of this script serves; the rest belong to child processes.
OURS = (AGENT_PORT, ENGINE_PROXY, EUGENE_PORT, GW_PROXY)
WORKBENCH = f"http://127.0.0.1:{WB_PORT}"
LOCAL = "mr4-local"

CLIENT_ID, CLIENT_SECRET, ADMIN_TOKEN = "mr4-client", "mr4-client-secret", "mr4-admin-token"
#: Made up, so only a model that was shown one can say it. Each file has its own words, so
#: a question has one right answer and the files do not compete for it.
CODE = "HERON-1180"  # survey.pdf, page 4, as first added
CODE_TWO = "OSPREY-2255"  # survey.pdf, page 4, once the file is added again
PANTRY = "QUOKKA-7359"  # notes.md
TIDE = "TERN-4417"  # tides.docx
CODE_CHAT = "KESTREL-9402"  # a file attached to one chat, page 9 of 30
BO_WORD = "PLOVER-3081"  # Bo's own file
CODE_WORD_ASKED = (
    "What is the code word for the lighthouse inspection? Reply with only the code word. /no_think"
)
PANTRY_ASKED = "What is the pantry padlock combination? Reply with only the combination. /no_think"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

FAILURES = 0
PASSES = 0


def say(text: str) -> None:
    print(f"\n== {text}", flush=True)


def check(claim: str, passed: bool, detail: object = "") -> bool:
    global FAILURES, PASSES
    if passed:
        PASSES += 1
    else:
        FAILURES += 1
    shown = f"  ({str(detail)[:400]})" if detail != "" and not passed else ""
    print(f"  {'PASS' if passed else 'FAIL'}  {claim}{shown}", flush=True)
    return passed


def note(text: str) -> None:
    print(f"  NOTE  {text}", flush=True)


def listening(port: int) -> bool:
    with socket.socket() as s:
        s.settimeout(0.3)
        return s.connect_ex(("127.0.0.1", port)) == 0


def stem(word: str) -> str:
    return word.split("-")[0]


def survey_pdf(word: str, *, page: int = 4, pages: int = 6, sentence: str | None = None) -> bytes:
    """A PDF of full pages, each about a passage long, with one sentence on one page only:
    by default the code word for the lighthouse inspection."""
    said = sentence or f"The code word for the lighthouse inspection is {word}."
    sheets = []
    for n in range(1, pages + 1):
        lines = [f"Page {n} of the harbour survey records ordinary observations in plain words."]
        lines += [
            f"Observation {n}.{i} describes the weather, the tide and the gulls." for i in range(21)
        ]
        if n == page:
            lines.insert(5, said)
        sheets.append("\n".join(lines))
    return docs.pdf(sheets)


def tides_docx(word: str) -> bytes:
    """A Word file, written with python-docx (which the unit checks' fixtures use too)."""
    import io

    import docx as word_files

    document = word_files.Document()
    document.add_heading("Tide tables", level=1)
    document.add_paragraph("The tables follow the gauge at the harbour mouth.")
    document.add_heading("Dawn readings", level=2)
    document.add_paragraph(f"The tide gauge reading at dawn on the spring tide is {word}.")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- #
# What each hop was sent, and what the engine said of its own cache
# --------------------------------------------------------------------------- #

#: Every chat request Workbench sent the gateway, and every one the engine received.
WB_SENT: list[dict[str, Any]] = []
ENGINE_BODIES: list[dict[str, Any]] = []
LOCK = threading.Lock()


class Quiet(ThreadingHTTPServer):
    daemon_threads = True

    def handle_error(self, request: Any, client_address: Any) -> None:
        # A probe that hangs up early is not news.
        pass


def make_proxy(upstream: int, sink: list[dict[str, Any]]) -> type[BaseHTTPRequestHandler]:
    """A pass-through to `127.0.0.1:upstream` that keeps the body of every chat
    request in `sink`: exactly what the next hop was sent."""

    class Proxy(BaseHTTPRequestHandler):
        def log_message(self, *args: Any) -> None:
            pass

        def _forward(self) -> None:
            length = int(self.headers.get("content-length") or 0)
            body = self.rfile.read(length) if length else b""
            if self.command == "POST" and self.path.split("?")[0] == "/v1/chat/completions":
                with LOCK:
                    sink.append(json.loads(body))
            hop = http.client.HTTPConnection("127.0.0.1", upstream, timeout=900)
            wanted = ("content-type", "accept", "authorization", "x-eugene-repetition-mode")
            headers = {k: v for k, v in self.headers.items() if k.lower() in wanted}
            hop.request(self.command, self.path, body=body or None, headers=headers)
            answer = hop.getresponse()
            self.send_response(answer.status)
            for key in ("Content-Type", "Cache-Control"):
                if value := answer.getheader(key):
                    self.send_header(key, value)
            self.end_headers()  # HTTP/1.0: the body ends when the connection does
            while chunk := answer.read(4096):
                self.wfile.write(chunk)
                self.wfile.flush()
            hop.close()

        do_GET = _forward
        do_POST = _forward
        do_PATCH = _forward
        do_DELETE = _forward

    return Proxy


class Agent(BaseHTTPRequestHandler):
    """Topology only: the gateway asks its agent which drivers there are."""

    def log_message(self, *args: Any) -> None:
        pass

    def do_GET(self) -> None:
        path = self.path.split("?")[0]
        if path == "/v1/components":
            body: Any = {
                "components": [
                    {
                        "name": "mr4-local-driver",
                        "kind": "inference-driver",
                        "url": f"http://127.0.0.1:{DRV_LOCAL}/",
                        "status": "running",
                    }
                ]
            }
        elif path == "/v1/runtimes":
            body = {"runtimes": []}
        elif path == "/v1/node":
            body = {"name": "mr4-node", "enrolled": False}
        else:
            body = {"status": "ok"}
        raw = json.dumps(body).encode()
        self.send_response(200)
        self.send_header("content-type", "application/json")
        self.send_header("content-length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)


def serve(handler: type[BaseHTTPRequestHandler], port: int) -> None:
    Quiet(("127.0.0.1", port), handler).serve_forever()


# --------------------------------------------------------------------------- #
# A stand-in for Eugene's sign-in (OpenID Connect, as the unit checks use)
# --------------------------------------------------------------------------- #

PEOPLE = {
    "operator": {"name": "Owner", "role": "operator", "username": "operator"},
    "p-ada": {"name": "Ada", "role": "member", "username": "ada"},
    "p-bo": {"name": "Bo", "role": "member", "username": "bo"},
}


def eugene_app() -> Any:
    key = RSAKey.generate_key(2048, parameters={"kid": "k1"})
    issuer = f"http://127.0.0.1:{EUGENE_PORT}/oidc"
    codes: dict[str, dict[str, Any]] = {}
    app = FastAPI()

    def id_token(sub: str, nonce: str | None) -> str:
        now = int(time.time())
        person = PEOPLE[sub]
        claims: dict[str, Any] = {
            "iss": issuer,
            "aud": CLIENT_ID,
            "sub": sub,
            "iat": now,
            "exp": now + 600,
            "name": person["name"],
            "preferred_username": person["username"],
            "eugene_role": person["role"],
            "auth_time": now,
        }
        if nonce is not None:
            claims["nonce"] = nonce
        return jwt.encode({"alg": "RS256", "typ": "JWT", "kid": "k1"}, claims, key)

    @app.get("/oidc/.well-known/openid-configuration")
    async def discovery() -> Any:
        return {
            "issuer": issuer,
            "authorization_endpoint": f"{issuer}/authorize",
            "token_endpoint": f"{issuer}/token",
            "jwks_uri": f"{issuer}/jwks",
            "revocation_endpoint": f"{issuer}/revoke",
        }

    @app.get("/oidc/jwks")
    async def jwks() -> Any:
        return {"keys": [key.as_dict(private=False)]}

    @app.get("/oidc/authorize")
    async def authorize(request: Request) -> Any:
        q = request.query_params
        code = secrets.token_urlsafe(16)
        codes[code] = {
            "sub": q["login_as"],
            "nonce": q["nonce"],
            "challenge": q["code_challenge"],
            "redirect_uri": q["redirect_uri"],
        }
        target = (
            q["redirect_uri"] + "?" + urlencode({"code": code, "state": q["state"], "iss": issuer})
        )
        return RedirectResponse(target, status_code=302)

    @app.post("/oidc/token")
    async def token(request: Request) -> Any:
        auth = request.headers.get("authorization", "")
        expected = base64.b64encode(f"{CLIENT_ID}:{CLIENT_SECRET}".encode()).decode()
        if auth != f"Basic {expected}":
            return JSONResponse({"error": "invalid_client"}, status_code=401)
        form = await request.form()
        if form["grant_type"] == "authorization_code":
            found = codes.pop(str(form["code"]), None)
            verifier = str(form.get("code_verifier", ""))
            challenge = (
                base64.urlsafe_b64encode(hashlib.sha256(verifier.encode()).digest())
                .rstrip(b"=")
                .decode()
            )
            if found is None or found["challenge"] != challenge:
                return JSONResponse({"error": "invalid_grant"}, status_code=400)
            return {
                "access_token": "at",
                "token_type": "Bearer",
                "expires_in": 3600,
                "id_token": id_token(found["sub"], found["nonce"]),
                "refresh_token": secrets.token_urlsafe(16),
            }
        return JSONResponse({"error": "unsupported_grant_type"}, status_code=400)

    @app.post("/oidc/revoke")
    async def revoke() -> Any:
        return {}

    return app


# --------------------------------------------------------------------------- #
# A person's browser
# --------------------------------------------------------------------------- #


class Person:
    """One person's browser: a cookie jar, the secret the page keeps, an origin."""

    def __init__(self, sub: str) -> None:
        self.sub = sub
        self.http = httpx.Client(base_url=WORKBENCH, timeout=240, trust_env=False)
        self.secret: str | None = None

    def sign_in(self) -> None:
        start = self.http.get("/signin")
        assert start.status_code == 302, start.text
        at_eugene = httpx.get(start.headers["location"] + "&login_as=" + self.sub, trust_env=False)
        assert at_eugene.status_code == 302, at_eugene.text
        back = urlsplit(at_eugene.headers["location"])
        done = self.http.get(back.path + "?" + back.query)
        location = done.headers.get("location", "")
        assert "#signin=" in location, f"sign-in failed: {location}"
        self.secret = location.split("#signin=", 1)[1]

    def call(self, method: str, path: str, **kw: Any) -> Any:
        headers = {"Origin": WORKBENCH, "X-Workbench-Secret": self.secret or ""}
        return self.http.request(method, path, headers=headers, **kw)

    def get(self, path: str, **kw: Any) -> Any:
        return self.call("GET", path, **kw)

    def post(self, path: str, **kw: Any) -> Any:
        return self.call("POST", path, **kw)

    def put(self, path: str, **kw: Any) -> Any:
        return self.call("PUT", path, **kw)

    def patch(self, path: str, **kw: Any) -> Any:
        return self.call("PATCH", path, **kw)

    def delete(self, path: str, **kw: Any) -> Any:
        return self.call("DELETE", path, **kw)

    def new_chat(self, model: str, **settings: Any) -> str:
        chat = self.post("/api/chats", json={"model": model}).json()["id"]
        if settings:
            changed = self.patch(f"/api/chats/{chat}", json={"settings": settings})
            assert changed.status_code == 200, changed.text
        return str(chat)

    # --- My documents ---------------------------------------------------

    def add(self, *named: tuple[str, bytes, str]) -> Any:
        return self.post("/api/documents", files=[("files", n) for n in named])

    def documents(self) -> dict[str, Any]:
        found = self.get("/api/documents")
        assert found.status_code == 200, found.text
        return dict(found.json())

    def wait_documents(
        self, seconds: float = 180, *, until: Any = None, every: float = 0.02
    ) -> dict[str, Any]:
        """The listing once every document has ended, or once `until(listing)` says."""
        end = time.perf_counter() + seconds
        listed = self.documents()
        while time.perf_counter() < end:
            listed = self.documents()
            if until is not None and until(listed):
                return listed
            if until is None and all(
                d["state"] in ("ready", "skipped", "failed") for d in listed["documents"]
            ):
                return listed
            time.sleep(every)
        raise AssertionError("the documents were not read in time")

    # --- chats ----------------------------------------------------------

    def wait_answer(self, chat: str, seconds: float = 300) -> dict[str, Any]:
        end = time.perf_counter() + seconds
        while time.perf_counter() < end:
            messages = self.get(f"/api/chats/{chat}").json()["messages"]
            if messages and messages[-1]["status"] != "running":
                return dict(messages[-1])
            time.sleep(0.3)
        raise AssertionError("the answer did not finish")

    def say(self, chat: str, content: str, **body: Any) -> Any:
        return self.post(f"/api/chats/{chat}/messages", json={"content": content, **body})


def newest(sink: list[dict[str, Any]], count: int, what: str) -> dict[str, Any]:
    """The newest request in `sink`, once it holds more than `count`."""
    end = time.time() + 60
    while time.time() < end:
        with LOCK:
            if len(sink) > count:
                return sink[-1]
        time.sleep(0.2)
    raise AssertionError(f"{what} was sent no request")


def text_of(message: dict[str, Any]) -> str:
    content = message["content"]
    return content if isinstance(content, str) else str(content[0]["text"])


BLOCK = re.compile(r"<documents-([0-9a-f]{8})>\n(.*?)\n</documents-\1>", re.DOTALL)
PASSAGE = re.compile(
    r'<document-([0-9a-f]{8}) n="(\d+)" name="([^"]*)"((?: [a-z]+="[^"]*")*)>'
    r"\n(.*?)\n</document-\1>",
    re.DOTALL,
)


def passages_in(message: dict[str, Any]) -> list[tuple[str, ...]]:
    return [m.groups() for m in PASSAGE.finditer(text_of(message))]


def person_key(sub: str) -> str:
    return hashlib.sha256(sub.encode("utf-8")).hexdigest()[:32]


# --------------------------------------------------------------------------- #
# The run
# --------------------------------------------------------------------------- #


def main() -> int:
    cast(Any, sys.stdout).reconfigure(encoding="utf-8", errors="replace")
    procs: list[subprocess.Popen[bytes]] = []
    work = Path(tempfile.mkdtemp(prefix="ep-mr4-"))
    logs = work / "logs"
    logs.mkdir()
    data = work / "workbench-data"
    wb_log_count = [0]

    def start(argv: list[str], log: str, env: dict[str, str]) -> subprocess.Popen[bytes]:
        # The child owns the handle for as long as it runs; teardown kills it.
        handle = open(logs / log, "w", encoding="utf-8")  # noqa: SIM115
        proc = subprocess.Popen(
            argv, stdin=subprocess.DEVNULL, stdout=handle, stderr=subprocess.STDOUT, env=env
        )
        procs.append(proc)
        return proc

    def wait_for(port: int, seconds: float = 120) -> bool:
        end = time.time() + seconds
        while time.time() < end:
            if listening(port):
                return True
            time.sleep(0.5)
        return False

    def kill(proc: subprocess.Popen[bytes]) -> None:
        subprocess.run(
            ["taskkill", "/PID", str(proc.pid), "/T", "/F"], capture_output=True, check=False
        )

    def counts() -> tuple[int, int]:
        with LOCK:
            return len(WB_SENT), len(ENGINE_BODIES)

    def index_path(sub: str) -> Path:
        return data / "knowledge" / person_key(sub) / "index.sqlite3"

    def knowledge_files(sub: str) -> list[Path]:
        root = data / "knowledge" / person_key(sub)
        return [p for p in root.rglob("*") if p.is_file()] if root.exists() else []

    server: Any = None
    wb: list[subprocess.Popen[bytes]] = []
    try:
        say("0. isolate this run")
        base = {k: v for k, v in os.environ.items() if not k.startswith("EUGENE_PLEXUS_")}
        for proxy in ("HTTP_PROXY", "HTTPS_PROXY", "ALL_PROXY", "NO_PROXY"):
            base.pop(proxy, None)
            base.pop(proxy.lower(), None)
        check("no ambient EUGENE_PLEXUS_* variable reaches a child", True)
        taken = [p for p in OWNED if listening(p)]
        if taken:
            check(f"ports {taken} are free", False)
            return 1
        for path in (
            WB_PY,
            GW_PY,
            DRV_PY,
            LLAMA,
            MODEL_FILE,
            WB / "src" / "eugene_plexus_workbench",
        ):
            if not path.exists():
                check(f"{path} exists", False)
                return 1
        check(
            f"ports {OWNED[0]}-{OWNED[-1]} free; Workbench from {WB}; llama-server "
            f"{LLAMA.parent.name}; {MODEL_FILE.name}",
            True,
        )

        say("1. the stack: engine, driver, gateway, Workbench, and three people")
        threading.Thread(target=serve, args=(Agent, AGENT_PORT), daemon=True).start()
        threading.Thread(
            target=serve, args=(make_proxy(ENGINE, ENGINE_BODIES), ENGINE_PROXY), daemon=True
        ).start()
        threading.Thread(
            target=serve, args=(make_proxy(GW_PORT, WB_SENT), GW_PROXY), daemon=True
        ).start()
        config = uvicorn.Config(eugene_app(), host="127.0.0.1", port=EUGENE_PORT, log_level="error")
        server = uvicorn.Server(config)
        threading.Thread(target=server.run, daemon=True).start()
        # CPU only, one slot (several would divide the window), a window of 8,192.
        engine = [str(LLAMA), "-m", str(MODEL_FILE), "--host", "127.0.0.1", "--port", str(ENGINE)]
        engine += ["-ngl", "0", "-t", "4", "-c", "8192", "-np", "1", "-b", "512", "-ub", "512"]
        start(
            engine,
            "llama.log",
            {**base, "CUDA_VISIBLE_DEVICES": "-1", "GGML_VK_VISIBLE_DEVICES": ""},
        )
        (work / "driver-local.yaml").write_text(
            f"provider: openai_compat_custom\nbaseUrl: http://127.0.0.1:{ENGINE_PROXY}\n"
            f"modelId: {LOCAL}\nbackendLocality: local\nrequestTimeoutSeconds: 600\n",
            encoding="utf-8",
            newline="\n",
        )
        start(
            [str(DRV_PY), "-m", "eugene_plexus_inference_driver"],
            "driver-local.log",
            {
                **base,
                "EUGENE_PLEXUS_DRIVER_CONFIG_FILE": str(work / "driver-local.yaml"),
                "EUGENE_PLEXUS_DRIVER_BIND_HOST": "127.0.0.1",
                "EUGENE_PLEXUS_DRIVER_BIND_PORT": str(DRV_LOCAL),
            },
        )
        (work / "gateway.yaml").write_text("routingRefreshSeconds: 3\n", encoding="utf-8")
        start(
            [str(GW_PY), "-m", "eugene_plexus_gateway"],
            "gateway.log",
            {
                **base,
                "EUGENE_PLEXUS_GATEWAY_CONFIG_FILE": str(work / "gateway.yaml"),
                "EUGENE_PLEXUS_GATEWAY_METRICS_FILE": str(work / "metrics.sqlite3"),
                "EUGENE_PLEXUS_GATEWAY_BIND_HOST": "127.0.0.1",
                "EUGENE_PLEXUS_GATEWAY_BIND_PORT": str(GW_PORT),
                "EUGENE_PLEXUS_GATEWAY_AGENT_URL": f"http://127.0.0.1:{AGENT_PORT}",
            },
        )
        for port in (ENGINE, DRV_LOCAL, GW_PORT, EUGENE_PORT):
            if not wait_for(port, 180):
                check(f"something listens on {port}; logs in {logs}", False)
                return 1
        ids: set[str] = set()
        known = False
        end = time.time() + 240
        while time.time() < end and not known:
            try:
                listed = httpx.get(
                    f"http://127.0.0.1:{GW_PORT}/v1/models", timeout=5, trust_env=False
                ).json()
                ids = {m["id"] for m in listed.get("data") or []}
                entry = next((m for m in listed.get("data") or [] if m["id"] == LOCAL), None)
                # The driver learns the engine's window once the model has loaded.
                known = (
                    entry is not None
                    and (entry.get("x_eugene_plexus") or {}).get("context_length") is not None
                )
            except (httpx.HTTPError, ValueError):
                pass
            time.sleep(1)
        if not check(f"the gateway routes {LOCAL} and knows its window", known, sorted(ids)):
            return 1

        static = work / "static"
        static.mkdir()
        (static / "index.html").write_text(
            "<!doctype html><title>Workbench</title>", encoding="utf-8"
        )
        (work / "app_key").write_text("mr4-app-key", encoding="utf-8")
        (work / "oidc_secret").write_text(CLIENT_SECRET, encoding="utf-8")

        def start_workbench() -> subprocess.Popen[bytes]:
            wb_log_count[0] += 1
            return start(
                [str(WB_PY), "-m", "eugene_plexus_workbench"],
                f"workbench-{wb_log_count[0]}.log",
                {
                    **base,
                    "PYTHONPATH": str(WB / "src"),
                    "EUGENE_PLEXUS_APP_BIND_HOST": "127.0.0.1",
                    "EUGENE_PLEXUS_APP_BIND_PORT": str(WB_PORT),
                    "EUGENE_PLEXUS_APP_DATA_DIR": str(data),
                    "EUGENE_PLEXUS_APP_KEY_FILE": str(work / "app_key"),
                    "EUGENE_PLEXUS_APP_ADMIN_TOKEN": ADMIN_TOKEN,
                    # Through the recording proxy: it is what Workbench sends the gateway.
                    "EUGENE_PLEXUS_APP_GATEWAY_URL": f"http://127.0.0.1:{GW_PROXY}",
                    "EUGENE_PLEXUS_APP_OIDC_ISSUER": f"http://127.0.0.1:{EUGENE_PORT}/oidc",
                    "EUGENE_PLEXUS_APP_OIDC_CLIENT_ID": CLIENT_ID,
                    "EUGENE_PLEXUS_APP_OIDC_SECRET_FILE": str(work / "oidc_secret"),
                    "EUGENE_PLEXUS_APP_STATIC_DIR": str(static),
                },
            )

        wb.append(start_workbench())
        if not wait_for(WB_PORT, 60):
            check(f"Workbench listens on {WB_PORT}; logs in {logs}", False)
            return 1
        ada, bo, owner = Person("p-ada"), Person("p-bo"), Person("operator")
        for person in (ada, bo, owner):
            person.sign_in()
        check("Ada, Bo and the owner sign in with Eugene", all(p.secret for p in (ada, bo, owner)))
        listing = {m["id"]: m for m in ada.get("/api/models").json()["models"]}
        window = listing[LOCAL]["contextLength"]
        note(f"{LOCAL}: window {window} tokens; locality {listing[LOCAL]['locality']}")
        if not check(
            "the model is a local one with a window of 8,192",
            window == 8192 and listing[LOCAL]["locality"] == "local",
            listing[LOCAL],
        ):
            return 1

        def settings(**values: Any) -> Any:
            return httpx.patch(
                f"{WORKBENCH}/v1/config",
                json=values,
                headers={"Authorization": f"Bearer {ADMIN_TOKEN}"},
                trust_env=False,
            )

        def reached_by(reply: dict[str, Any], before_engine: int) -> dict[str, Any]:
            """What the engine received for this answer; nothing when no model was asked."""
            if reply["status"] != "done" or reply.get("finish") == "documents_only":
                return {}
            return newest(ENGINE_BODIES, before_engine, "the engine")

        def ask(
            person: Person,
            chat: str,
            question: str,
            *,
            word: str | None = None,
            tries: int = 3,
            **body: Any,
        ) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
            """Ask, wait for the real model, and ask again (a small model is not perfect)
            until its answer holds `word`. The answer, what Workbench sent the gateway, and
            what reached the engine, for the last try."""
            before_sent, before_engine = counts()
            sent = person.say(chat, question, **body)
            assert sent.status_code == 201, sent.text
            reply = person.wait_answer(chat)
            request = newest(WB_SENT, before_sent, "the gateway")
            reached = reached_by(reply, before_engine)
            tried = 1
            while word is not None and stem(word) not in reply["content"].upper() and tried < tries:
                tried += 1
                before_sent, before_engine = counts()
                person.post(f"/api/chats/{chat}/retry")
                reply = person.wait_answer(chat)
                request = newest(WB_SENT, before_sent, "the gateway")
                reached = reached_by(reply, before_engine)
            return reply, request, reached

        def holds(sub: str, word: str) -> list[str]:
            """Every file of a person's own folder that holds this word, as it was written
            or as the full-text index keeps it (in lower case, a word at a time): the index,
            its log, an original, the text read from it."""
            needles = (word.lower().encode(), stem(word).lower().encode())
            return [
                p.name
                for p in knowledge_files(sub)
                if not p.name.endswith("-shm") and any(n in p.read_bytes().lower() for n in needles)
            ]

        # ------------------------------------------------------------------ #
        say("2. three files into My documents: each file's own state, then the counts")
        notes_md = (
            "# Keeper's notes\n\nThe keeper writes about the garden and the kettle.\n\n"
            f"## The pantry\n\nThe pantry padlock combination is {PANTRY}, known to nobody but "
            "the old keeper.\n"
        ).encode()
        made = ada.add(
            ("survey.pdf", survey_pdf(CODE), "application/pdf"),
            ("notes.md", notes_md, "text/markdown"),
            ("tides.docx", tides_docx(TIDE), DOCX_TYPE),
        )
        body = made.json()
        check(
            "three files are kept at once, each queued, and nothing is refused",
            made.status_code == 201
            and body["refused"] == []
            and [d["name"] for d in body["documents"]] == ["survey.pdf", "notes.md", "tides.docx"]
            and {d["state"] for d in body["documents"]} == {"queued"},
            made.text[:300],
        )
        seen: dict[str, list[str]] = {d["name"]: [d["state"]] for d in body["documents"]}
        end = time.perf_counter() + 180
        listed = ada.documents()
        while time.perf_counter() < end:
            listed = ada.documents()
            for d in listed["documents"]:
                if seen[d["name"]][-1] != d["state"]:
                    seen[d["name"]].append(d["state"])
            if all(d["state"] in ("ready", "skipped", "failed") for d in listed["documents"]):
                break
            time.sleep(0.01)
        for name, states in seen.items():
            note(f"{name}: " + " -> ".join(states))
        got = {d["name"]: d for d in listed["documents"]}
        check(
            "every file went through its own states (queued, then reading) to ready",
            all(s[0] == "queued" and s[-1] == "ready" and "reading" in s for s in seen.values()),
            seen,
        )
        check(
            "each file's counts: the PDF its six pages, the Markdown its lines, every one its "
            "passages, size and kind",
            got["survey.pdf"]["pages"] == 6
            and got["survey.pdf"]["label"] == "PDF"
            and got["notes.md"]["label"] == "Markdown"
            and (got["notes.md"]["lines"] or 0) >= 5
            and got["tides.docx"]["label"] == "Word"
            and all(d["passages"] >= 1 and d["tokens"] and d["size"] > 0 for d in got.values()),
            got,
        )
        totals = listed["totals"]
        check(
            "the totals add up, and Workbench says it searches by words only",
            totals["documents"] == 3
            and totals["byState"]["ready"] == 3
            and totals["passages"] == sum(d["passages"] for d in got.values())
            and listed["search"]["headline"] == "Searching by words only.",
            totals,
        )
        note(
            f"{totals['passages']} passages in {totals['documents']} files; "
            f"{got['survey.pdf']['passages']} from the six-page PDF"
        )
        by_own_words = {
            "notes.md": ("pantry padlock combination", "The pantry"),
            "tides.docx": ("gauge reading at dawn on the spring", "Dawn readings"),
        }
        for name, (asked, heading) in by_own_words.items():
            looked = ada.post("/api/documents/test", json={"question": asked}).json()
            top = looked["passages"][0] if looked["passages"] else {}
            check(
                f"{name} is found by its own words, under its own heading, and by no other file",
                looked["status"] == "found"
                and {p["name"] for p in looked["passages"]} == {name}
                and heading in (top.get("heading") or ""),
                top or looked,
            )

        # ------------------------------------------------------------------ #
        say(
            "3. a question whose answer is a code word on page 4: the model gives it, with its page"
        )
        before = counts()[0]
        tried = ada.post("/api/documents/test", json={"question": CODE_WORD_ASKED})
        found = tried.json()
        top = found["passages"][0] if found["passages"] else {}
        check(
            "Test a question shows the passages with file, page and why, and asks no model",
            tried.status_code == 200
            and counts()[0] == before
            and found["status"] == "found"
            and top.get("name") == "survey.pdf"
            and (top.get("pageStart"), top.get("pageEnd")) == (4, 4)
            and "lighthouse" in top["why"]["terms"]
            and CODE in top["text"],
            top or found,
        )
        survey_id = got["survey.pdf"]["id"]
        shown = ada.get(f"/api/documents/{survey_id}/text")
        check(
            "Show indexed text is the text with its pages marked",
            shown.status_code == 200 and "[page 4]" in shown.text and CODE in shown.text,
        )
        # The switch is off until the person turns it on: a chat that has not is sent
        # nothing from the documents, and says nothing about them.
        off_chat = ada.new_chat(LOCAL, temperature=0, maxTokens=100)
        reply, request, _ = ask(ada, off_chat, CODE_WORD_ASKED)
        check(
            "Documents are off in a new chat: nothing from the documents is sent, and the answer "
            "says nothing about searching",
            ada.get(f"/api/chats/{off_chat}").json()["chat"]["settings"]["documents"] is None
            and "<documents-" not in json.dumps(request["messages"])
            and CODE not in json.dumps(request["messages"])
            and reply["retrieval"] is None
            and reply["sources"] == [],
            reply["retrieval"],
        )
        chat = ada.new_chat(LOCAL, temperature=0, maxTokens=300, documents="on")
        chip = ada.get(f"/api/chats/{chat}/documents").json()
        check(
            "the chat's chip says what its next request will do, from the same decision the "
            "request makes: it searches My documents, in a window of 8,192, within 2,457 tokens",
            chip["switch"] == "on"
            and chip["sending"] is True
            and chip["searching"] == {"library": True, "chat": False}
            and chip["reason"] is None
            and chip["model"]["locality"] == "local"
            and chip["counts"]["library"]["documents"] == 3
            and chip["budget"] == {"limit": 2457, "window": 8192},
            chip,
        )
        reply, request, reached = ask(ada, chat, CODE_WORD_ASKED, word=CODE)
        latest = request["messages"][-1]
        found_passages = passages_in(latest)
        check(
            "Workbench's request is ONE user message: the passages block first, the question last",
            [m["role"] for m in request["messages"]] == ["user"]
            and isinstance(latest["content"], str)
            and BLOCK.search(latest["content"]) is not None
            and latest["content"].startswith("<documents-")
            and latest["content"].endswith("\n\n" + CODE_WORD_ASKED),
            [m["role"] for m in request["messages"]],
        )
        check(
            "the first passage is survey.pdf page 4 with the code word, numbered 1",
            bool(found_passages)
            and found_passages[0][1] == "1"
            and found_passages[0][2] == "survey.pdf"
            and found_passages[0][0] == BLOCK.search(latest["content"]).group(1)  # type: ignore[union-attr]
            and 'pages="4"' in found_passages[0][3]
            and CODE in found_passages[0][4],
            found_passages[:1],
        )
        check(
            "and the same passage reached the real engine, through the gateway and the driver",
            bool(reached)
            and CODE in json.dumps(reached["messages"])
            and "<documents-" in json.dumps(reached["messages"]),
        )
        check(
            "the real model answers with the code word it was given",
            stem(CODE) in reply["content"].upper(),
            reply["content"][:150],
        )
        note(f"answer: {reply['content'][:100]!r}")
        passages = [s for s in reply["sources"] if s.get("passage")]
        check(
            "the Sources entry names the file and the page, how it was found, and its rank",
            bool(passages)
            and passages[0]["kind"] == "document"
            and passages[0]["name"] == "survey.pdf"
            and (passages[0]["page_start"], passages[0]["page_end"]) == (4, 4)
            and passages[0]["rank"] == 1
            and passages[0]["via"] == "words"
            and CODE in passages[0]["snippet"]
            and len(passages[0]["snippet"]) <= 240,
            passages[:1],
        )
        check(
            "the answer keeps what searching came to, and the stored message is the person's words",
            reply["retrieval"]["status"] == "used"
            and reply["retrieval"]["scope"] == {"library": True, "chat": False}
            and reply["retrieval"]["passages"] == len(passages)
            and ada.get(f"/api/chats/{chat}").json()["messages"][0]["content"] == CODE_WORD_ASKED,
            reply["retrieval"],
        )
        spent = reply["retrieval"]
        check(
            "and the budget it was held to: at most min(3,000 tokens, 30% of the window), "
            "the window the gateway reported, and what was sent is inside it",
            spent["budget"]["window"] == 8192
            and spent["budget"]["limit"] == 2457
            and spent["budget"]["allowed"] <= spent["budget"]["limit"]
            and 0 < spent["tokens"] <= spent["budget"]["allowed"]
            and spent["searched"]["documents"] == 3
            and spent["found"] >= spent["passages"] >= 1
            and spent["mode"] == "words",
            spent,
        )
        # The next turn: the passages are not sent again, and are looked for again.
        _reply_two, request2, _ = ask(
            ada, chat, "Thank you. And what does the lighthouse inspection page say? /no_think"
        )
        sent_two = request2["messages"]
        check(
            "the next request does not replay them: the earlier turn is the person's words, and "
            "only the latest message carries a block, under a suffix of its own",
            [m["role"] for m in sent_two] == ["user", "assistant", "user"]
            and text_of(sent_two[0]) == CODE_WORD_ASKED
            and "<documents-" not in json.dumps(sent_two[:2])
            and "<documents-" in text_of(sent_two[2])
            and BLOCK.search(text_of(sent_two[2])).group(1)  # type: ignore[union-attr]
            != BLOCK.search(latest["content"]).group(1),  # type: ignore[union-attr]
            [m["role"] for m in sent_two],
        )
        pantry_chat = ada.new_chat(LOCAL, temperature=0, maxTokens=300, documents="on")
        pantry_before = ask(ada, pantry_chat, PANTRY_ASKED, word=PANTRY)
        check(
            "and a question about another file is answered from that file, by the real model",
            PANTRY in json.dumps(pantry_before[1]["messages"])
            and CODE not in text_of(pantry_before[1]["messages"][-1])
            and stem(PANTRY) in pantry_before[0]["content"].upper(),
            pantry_before[0]["content"][:120],
        )

        # ------------------------------------------------------------------ #
        say("4. edit (add the file again) or delete: the old words are gone from answers at once")
        again = ada.add(("survey.pdf", survey_pdf(CODE_TWO), "application/pdf"))
        [replaced] = again.json()["documents"]
        check(
            "the same name replaces the old file: its old id is named, the new one is queued, "
            "and the old one is gone from the list at once",
            again.status_code == 201
            and replaced["replaced"] == [survey_id]
            and replaced["state"] == "queued"
            and [d["id"] for d in ada.documents()["documents"] if d["name"] == "survey.pdf"]
            == [replaced["id"]]
            and ada.get(f"/api/documents/{survey_id}").status_code == 404,
            replaced,
        )
        edit_chat = ada.new_chat(LOCAL, temperature=0, maxTokens=300, documents="on")
        reply, request, reached = ask(ada, edit_chat, CODE_WORD_ASKED)
        check(
            "asked straight away, no hop was sent the old code word",
            CODE not in json.dumps(request["messages"]) and CODE not in json.dumps(reached),
        )
        ada.wait_documents()
        reply, request, reached = ask(ada, edit_chat, CODE_WORD_ASKED, word=CODE_TWO)
        check(
            "once the new file is read, the new word is what is sent and what the model says",
            CODE_TWO in json.dumps(request["messages"])
            and CODE not in json.dumps(request["messages"])
            and stem(CODE_TWO) in reply["content"].upper(),
            reply["content"][:150],
        )
        check(
            "and nothing of the old file is on the disk: no row, no file by its id, and its words "
            "in no file of the person's folder (the index, its log, the original, the text)",
            not any(survey_id in p.name for p in knowledge_files("p-ada"))
            and holds("p-ada", CODE) == [],
            holds("p-ada", CODE),
        )
        notes_id = got["notes.md"]["id"]
        gone = ada.delete(f"/api/documents/{notes_id}")
        after_delete = ada.new_chat(LOCAL, temperature=0, maxTokens=300, documents="on")
        reply, request, reached = ask(ada, after_delete, PANTRY_ASKED)
        check(
            "deleted: 204, gone from the list, its routes 404, its words are in no later request "
            "or answer, and nothing of it is on the disk",
            gone.status_code == 204
            and notes_id not in {d["id"] for d in ada.documents()["documents"]}
            and ada.get(f"/api/documents/{notes_id}/text").status_code == 404
            and ada.get(f"/api/documents/{notes_id}").status_code == 404
            and reply["retrieval"]["status"] == "nothing_matched"
            and PANTRY not in json.dumps(request["messages"])
            and PANTRY not in json.dumps(reached)
            and stem(PANTRY) not in reply["content"].upper()
            and not any(notes_id in p.name for p in knowledge_files("p-ada"))
            and holds("p-ada", PANTRY) == [],
            (reply["content"][:100], reply["retrieval"], holds("p-ada", PANTRY)),
        )
        again_asked = ask(ada, pantry_chat, PANTRY_ASKED)
        check(
            "and the chat that was answered from it before is, asked again, told nothing matched: "
            "its passages are looked for again each time, never kept in the chat",
            again_asked[0]["retrieval"]["status"] == "nothing_matched"
            and "<documents-" not in text_of(again_asked[1]["messages"][-1]),
            again_asked[0]["retrieval"],
        )

        # ------------------------------------------------------------------ #
        say("5. Workbench is killed in the middle of reading twelve files and starts again")
        many = [
            (
                f"batch{n:02d}.pdf",
                survey_pdf(
                    f"BATCH-{n:04d}",
                    page=3,
                    pages=5,
                    sentence=f"The manifest number for this batch is BATCH-{n:04d}.",
                ),
                "application/pdf",
            )
            for n in range(12)
        ]
        before_names = {d["name"] for d in ada.documents()["documents"]}
        made = ada.add(*many)
        check(
            "twelve files are kept at once, queued",
            made.status_code == 201 and len(made.json()["documents"]) == 12,
            made.text[:200],
        )

        def mid_way(listing_now: dict[str, Any]) -> bool:
            batch_now = [d for d in listing_now["documents"] if d["name"].startswith("batch")]
            return sum(d["state"] == "ready" for d in batch_now) >= 3 and any(
                d["state"] in ("queued", "reading") for d in batch_now
            )

        listed = ada.wait_documents(until=mid_way, every=0.005)
        ready_before = {
            d["name"]
            for d in listed["documents"]
            if d["name"].startswith("batch") and d["state"] == "ready"
        }
        kill(wb[-1])
        end = time.time() + 30
        while listening(WB_PORT) and time.time() < end:
            time.sleep(0.2)
        check("Workbench is killed", not listening(WB_PORT))
        with contextlib.closing(sqlite3.connect(index_path("p-ada"))) as db:
            at_death = {
                name: state for name, state in db.execute("SELECT name, status FROM documents")
            }
        cut_off = sorted(name for name, state in at_death.items() if state == "reading")
        waiting = sorted(name for name, state in at_death.items() if state == "queued")
        note(
            f"at death: {sum(s == 'ready' for s in at_death.values())} ready, "
            f"{len(cut_off)} being read ({', '.join(cut_off)}), {len(waiting)} waiting"
        )
        check(
            "its index says what was true when it died: some read, one cut off in the middle, "
            "the rest waiting",
            any(s == "ready" for s in at_death.values())
            and len(cut_off) == 1
            and len(waiting) >= 1,
            at_death,
        )
        wb.append(start_workbench())
        if not wait_for(WB_PORT, 60):
            check("Workbench starts again", False)
            return 1
        resumed = ada.wait_documents(seconds=240)
        names = [d["name"] for d in resumed["documents"]]
        batch = [d for d in resumed["documents"] if d["name"].startswith("batch")]
        check(
            "it resumes by itself: all twelve are ready, and none is left saying reading",
            len(batch) == 12 and all(d["state"] == "ready" for d in batch),
            [(d["name"], d["state"]) for d in batch],
        )
        with contextlib.closing(sqlite3.connect(index_path("p-ada"))) as db:
            attempts = dict(db.execute("SELECT name, attempt FROM documents"))
            dupes = db.execute(
                "SELECT doc_id, ord, COUNT(*) FROM chunks GROUP BY doc_id, ord HAVING COUNT(*) > 1"
            ).fetchall()
            problems = db.execute("PRAGMA integrity_check").fetchall()
        check(
            "each file was read once, and only the one the stop cut off was read again",
            len(names) == len(set(names)) == len(before_names) + 12
            and all(attempts[name] == (1 if name in cut_off else 0) for name in attempts),
            {n: a for n, a in attempts.items() if a},
        )
        check(
            "what was read before the stop was not lost, no passage is in the index twice, "
            "nothing half-written is left, and SQLite finds the file sound",
            ready_before <= {d["name"] for d in batch if d["state"] == "ready"}
            and all(d["passages"] >= 5 for d in batch)
            and resumed["totals"]["passages"] == sum(d["passages"] for d in resumed["documents"])
            and dupes == []
            and problems == [("ok",)]
            and not [p for p in knowledge_files("p-ada") if p.name.endswith(".part")],
            (dupes, problems),
        )
        by_manifest = []
        for n in range(12):
            looked = ada.post(
                "/api/documents/test", json={"question": f"manifest number BATCH-{n:04d}"}
            ).json()
            top = looked["passages"][0] if looked["passages"] else {}
            by_manifest.append(top.get("name") == f"batch{n:02d}.pdf" and top.get("pageStart") == 3)
        check(
            "every one of the twelve is found by its own manifest number, on page 3, by the "
            "restarted Workbench (the one that was cut off too)",
            all(by_manifest),
            by_manifest,
        )
        restart_chat = ada.new_chat(LOCAL, temperature=0, maxTokens=300, documents="on")
        reply, request, reached = ask(ada, restart_chat, CODE_WORD_ASKED, word=CODE_TWO)
        check(
            "and the restarted Workbench still answers from what it kept, with the real model",
            stem(CODE_TWO) in reply["content"].upper()
            and CODE_TWO in json.dumps(request["messages"]),
            reply["content"][:120],
        )
        for item in batch:
            ada.delete(f"/api/documents/{item['id']}")
        ada.wait_documents()

        # ------------------------------------------------------------------ #
        say("6. a PDF with no text, and a locked one: skipped, each with its reason")
        ada.add(
            ("scan.pdf", docs.pdf_pictures(3), "application/pdf"),
            ("locked.pdf", docs.pdf_encrypted(), "application/pdf"),
            ("empty.txt", b"  \n", "text/plain"),
        )
        listed = ada.wait_documents()
        skipped = {d["name"]: d for d in listed["documents"]}
        check(
            "the scan is skipped with its reason in words, and still says how many pages it has",
            skipped["scan.pdf"]["state"] == "skipped"
            and skipped["scan.pdf"]["code"] == "no_text"
            and skipped["scan.pdf"]["message"]
            == "Skipped: this PDF has no text, only pictures of pages. "
            "Workbench cannot read scanned pages yet."
            and skipped["scan.pdf"]["pages"] == 3
            and skipped["scan.pdf"]["passages"] == 0
            and skipped["scan.pdf"]["searchable"] is False,
            skipped["scan.pdf"],
        )
        check(
            "the locked file and the empty one are skipped in their own words",
            skipped["locked.pdf"]["message"] == "Skipped: this file is password-protected."
            and skipped["empty.txt"]["message"] == "Skipped: this file has no text in it.",
            (skipped["locked.pdf"]["message"], skipped["empty.txt"]["message"]),
        )
        check(
            "and the other files are unharmed: the totals count each state, and a skipped file "
            "has no text to show",
            listed["totals"]["byState"]["skipped"] == 3
            and listed["totals"]["byState"]["ready"] == 2
            and ada.get(f"/api/documents/{skipped['scan.pdf']['id']}/text").status_code == 409,
            listed["totals"],
        )
        for name in ("scan.pdf", "locked.pdf", "empty.txt"):
            ada.delete(f"/api/documents/{skipped[name]['id']}")
        tides = next(d for d in listed["documents"] if d["name"] == "tides.docx")
        reread = ada.post(f"/api/documents/{tides['id']}/reread")
        after = ada.wait_documents()
        tides_after = next(d for d in after["documents"] if d["name"] == "tides.docx")
        check(
            "Re-read reads a file again from the person's own copy: its old passages go at once "
            "(queued, with none), and it comes back as it was",
            reread.status_code == 200
            and reread.json()["state"] == "queued"
            and reread.json()["passages"] == 0
            and tides_after["state"] == "ready"
            and tides_after["id"] == tides["id"]
            and tides_after["passages"] == tides["passages"],
            (reread.text[:200], tides_after),
        )

        # ------------------------------------------------------------------ #
        say("7. nothing matched is one thing, a search that cannot run is another")
        quiet = ada.new_chat(LOCAL, temperature=0, maxTokens=100, documents="on")
        reply, request, _ = ask(
            ada, quiet, "What is the airspeed velocity of an unladen swallow? /no_think"
        )
        check(
            "a question the files do not answer says nothing matched, as a normal result, and "
            "the model is asked without passages",
            reply["retrieval"]["status"] == "nothing_matched"
            and reply["retrieval"]["message"] == "Nothing in your documents matched."
            and "<documents-" not in json.dumps(request["messages"])
            and reply["status"] == "done"
            and reply["sources"] == [],
            reply["retrieval"],
        )
        with contextlib.closing(sqlite3.connect(index_path("p-ada"))) as db:
            db.execute("DROP TABLE chunks_fts")
            db.commit()
        broken_chat = ada.new_chat(LOCAL, temperature=0, maxTokens=100, documents="on")
        reply, request, _ = ask(ada, broken_chat, CODE_WORD_ASKED)
        got_r = reply["retrieval"]
        check(
            "with the search index damaged underneath it, the answer goes on without passages and "
            "says the search failed, with the cause, and not that nothing matched",
            reply["status"] == "done"
            and got_r["status"] == "failed"
            and got_r["code"] == "index_damaged"
            and "The search index for your documents is damaged (SQLite said:" in got_r["message"]
            and "Your files are safe." in got_r["message"]
            and "Nothing in your documents matched" not in got_r["message"]
            and "<documents-" not in json.dumps(request["messages"]),
            got_r,
        )
        note(got_r["message"])
        broken = ada.post("/api/documents/test", json={"question": CODE_WORD_ASKED})
        check(
            "Test a question is a 503 with the same words and says Rebuild can fix it",
            broken.status_code == 503
            and broken.json()["detail"]["code"] == "index_damaged"
            and broken.json()["detail"]["canRebuild"] is True
            and "Your files are safe." in broken.json()["detail"]["message"],
            broken.text[:300],
        )
        still_listed = ada.get("/api/documents")
        check(
            "the list of documents still opens, with every file in it, and its totals",
            still_listed.status_code == 200
            and {d["name"] for d in still_listed.json()["documents"]}
            == {"survey.pdf", "tides.docx"},
            still_listed.text[:200],
        )
        rebuilt = ada.post("/api/documents/rebuild")
        listed = ada.wait_documents()
        check(
            "Rebuild makes the index again from the person's own files: nothing is lost",
            rebuilt.status_code == 200
            and rebuilt.json() == {"documents": 2, "left": 0}
            and {d["name"]: d["state"] for d in listed["documents"]}
            == {"survey.pdf": "ready", "tides.docx": "ready"},
            (rebuilt.text, [(d["name"], d["state"]) for d in listed["documents"]]),
        )
        rebuilt_chat = ada.new_chat(LOCAL, temperature=0, maxTokens=300, documents="on")
        reply, request, _ = ask(ada, rebuilt_chat, CODE_WORD_ASKED, word=CODE_TWO)
        check(
            "and the question is answered from the rebuilt index",
            reply["retrieval"]["status"] == "used" and stem(CODE_TWO) in reply["content"].upper(),
            reply["content"][:120],
        )

        # ------------------------------------------------------------------ #
        say("8. documents only: Workbench answers in its own words and asks no model")
        strict = ada.new_chat(
            LOCAL, temperature=0, maxTokens=200, documents="on", documentsOnly=True
        )
        before_sent, before_engine = counts()
        sent = ada.say(strict, "What is the airspeed velocity of an unladen swallow? /no_think")
        reply = ada.wait_answer(strict)
        check(
            "nothing matched in a documents-only chat: Workbench's own words, no model asked at "
            "any hop",
            sent.status_code == 201
            and reply["content"] == "I could not find that in your documents."
            and reply["finish"] == "documents_only"
            and reply["model"] is None
            and reply["retrieval"]["answeredBy"] == "workbench"
            and counts() == (before_sent, before_engine),
            (reply["content"], counts(), (before_sent, before_engine)),
        )
        ada.post(f"/api/chats/{strict}/retry", json={"askAnyway": True})
        reply = ada.wait_answer(strict)
        check(
            "Ask the model anyway asks it, and the answer says it was asked anyway",
            reply["finish"] != "documents_only"
            and reply["retrieval"]["askedAnyway"] is True
            and counts()[1] == before_engine + 1,
            reply,
        )
        reply, request, reached = ask(ada, strict, CODE_WORD_ASKED, word=CODE_TWO)
        check(
            "with a match, the model is told the passages are all there is",
            "Answer only from these passages and cite by number, like [1]."
            in json.dumps(request["messages"])
            and stem(CODE_TWO) in reply["content"].upper(),
            reply["content"][:120],
        )

        # The working indicator: "Searching your documents".
        progress_chat = ada.new_chat(LOCAL, temperature=0, maxTokens=100, documents="on")
        events: list[dict[str, Any]] = []
        watching = threading.Event()

        def watch() -> None:
            with ada.http.stream(
                "GET",
                f"/api/chats/{progress_chat}/events",
                headers={"Origin": WORKBENCH, "X-Workbench-Secret": ada.secret or ""},
                timeout=120,
            ) as response:
                for line in response.iter_lines():
                    if line.startswith(":"):
                        watching.set()
                    elif line.startswith("data: "):
                        event = json.loads(line[6:])
                        events.append(event)
                        if event.get("type") == "done":
                            return

        watcher = threading.Thread(target=watch, daemon=True)
        watcher.start()
        watching.wait(10)
        ada.say(progress_chat, CODE_WORD_ASKED)
        watcher.join(timeout=120)
        kinds = [e.get("type") for e in events]
        progress = [e["progress"] for e in events if e.get("type") == "progress"]
        check(
            "the working indicator is told it is searching the documents, before any word arrives",
            {"stage": "documents", "phase": "searching"} in progress
            and kinds.index("progress") < kinds.index("done"),
            kinds[:8],
        )

        # ------------------------------------------------------------------ #
        say(
            "9. a file attached to a chat and searched instead: the model answers from it, "
            "and it goes with the chat"
        )
        attach_chat = ada.new_chat(LOCAL, temperature=0, maxTokens=300)
        long_file = survey_pdf(CODE_CHAT, page=9, pages=30)
        up = ada.post(
            f"/api/chats/{attach_chat}/files",
            files={"file": ("long-survey.pdf", long_file, "application/pdf")},
        ).json()

        def own_file() -> dict[str, Any]:
            listing_now = ada.get(f"/api/chats/{attach_chat}/files").json()["files"]
            return dict(next(f for f in listing_now if f["id"] == up["id"]))

        end = time.time() + 120
        attached: dict[str, Any] = own_file()
        while time.time() < end and attached["reading"]["state"] == "reading":
            time.sleep(0.2)
            attached = own_file()
        offers = [o["action"] for o in (attached["decision"] or {}).get("offers", [])]
        check(
            "too long to read whole in this window, and Search it instead is offered with the "
            "other ways",
            attached["decision"]["kind"] in ("too_long", "choose")
            and attached["decision"]["send"] is False
            and offers[-1] == "search",
            attached["decision"],
        )
        chosen = ada.put(
            f"/api/chats/{attach_chat}/files/{up['id']}/reading", json={"mode": "search"}
        )
        end = time.time() + 120
        state: dict[str, Any] = chosen.json().get("search") or {}
        while time.time() < end and state.get("status") not in (
            "searchable",
            "ready",
            "skipped",
            "failed",
        ):
            time.sleep(0.2)
            state = own_file()["search"] or {}
        check(
            "chosen: it becomes a document of this chat, read in the background, and not one of "
            "My documents",
            chosen.status_code == 200
            and state.get("status") in ("searchable", "ready")
            and state.get("passages", 0) >= 25
            and "long-survey.pdf" not in {d["name"] for d in ada.documents()["documents"]},
            state,
        )
        reply, request, reached = ask(
            ada, attach_chat, CODE_WORD_ASKED, word=CODE_CHAT, attachments=[up["id"]]
        )
        latest = request["messages"][-1]
        found_passages = passages_in(latest)
        check(
            "the model answers from the file, and the file itself was not read into the message",
            stem(CODE_CHAT) in reply["content"].upper()
            and bool(found_passages)
            and found_passages[0][2] == "long-survey.pdf"
            and 'pages="9"' in found_passages[0][3]
            and CODE_CHAT in found_passages[0][4]
            and 'type="pdf"' not in text_of(latest)
            and "[page 1]" not in text_of(latest),
            (reply["content"][:100], found_passages[:1]),
        )
        passages = [s for s in reply["sources"] if s.get("passage")]
        check(
            "its Sources entry names the file and page 9, as the chat's own",
            bool(passages)
            and passages[0]["name"] == "long-survey.pdf"
            and passages[0]["scope"] == "chat"
            and (passages[0]["page_start"], passages[0]["page_end"]) == (9, 9),
            passages[:1],
        )
        other = ada.new_chat(LOCAL, temperature=0, maxTokens=100, documents="on")
        _, other_request, _ = ask(ada, other, CODE_WORD_ASKED)
        check(
            "no other chat is sent it: a chat with Documents on is sent My documents only",
            CODE_CHAT not in json.dumps(other_request["messages"])
            and CODE_TWO in json.dumps(other_request["messages"]),
        )
        with contextlib.closing(sqlite3.connect(index_path("p-ada"))) as db:
            in_chat = db.execute(
                "SELECT COUNT(*) FROM documents WHERE scope = ?", (f"chat:{attach_chat}",)
            ).fetchone()[0]
        check("while the chat lives, its document is in the person's index", in_chat == 1, in_chat)
        deleted = ada.delete(f"/api/chats/{attach_chat}")
        with contextlib.closing(sqlite3.connect(index_path("p-ada"))) as db:
            left = db.execute(
                "SELECT COUNT(*) FROM documents WHERE scope = ?", (f"chat:{attach_chat}",)
            ).fetchone()[0]
            hits = db.execute(
                "SELECT rowid FROM chunks_fts WHERE chunks_fts MATCH ?", (stem(CODE_CHAT),)
            ).fetchall()
        check(
            "deleting the chat deletes its document: its row, its passages, its search entries, "
            "its files, and its words are in no file of the folder",
            deleted.status_code == 204
            and left == 0
            and hits == []
            and holds("p-ada", CODE_CHAT) == []
            and ada.get(f"/api/chats/{attach_chat}").status_code == 404,
            (left, hits, holds("p-ada", CODE_CHAT)),
        )

        # ------------------------------------------------------------------ #
        say("10. two people on every endpoint, and the owner")
        target = next(d for d in ada.documents()["documents"] if d["name"] == "survey.pdf")
        bo.add(
            (
                "bo.md",
                f"# Bo's drawer\n\nBo keeps the lighthouse inspection code word {BO_WORD} in his "
                "drawer.\n".encode(),
                "text/markdown",
            )
        )
        bo.wait_documents()
        mine = bo.documents()["documents"][0]
        missing = "doesnotexist1"
        same = True
        detail: list[Any] = []
        for method, route in (
            ("get", "/api/documents/{id}"),
            ("delete", "/api/documents/{id}"),
            ("post", "/api/documents/{id}/reread"),
            ("get", "/api/documents/{id}/text"),
            ("get", "/api/documents/{id}/original"),
            ("get", "/api/documents/{id}/passages?ids=1,2"),
        ):
            theirs = getattr(bo, method)(route.format(id=target["id"]))
            nothing = getattr(bo, method)(route.format(id=missing))
            ok = (
                theirs.status_code == nothing.status_code == 404 and theirs.json() == nothing.json()
            )
            same = same and ok
            detail.append((method, route, theirs.status_code, nothing.status_code))
        check(
            "Bo gets the same 404 for Ada's document as for one that does not exist, on every "
            "route that takes a document",
            same,
            detail,
        )
        check(
            "and Ada's document is still there, whole, after all of that",
            ada.get(f"/api/documents/{target['id']}").json()["state"] == "ready"
            and CODE_TWO in ada.get(f"/api/documents/{target['id']}/text").text,
        )
        bo_chat = bo.new_chat(LOCAL, temperature=0, maxTokens=100, documents="on")
        check(
            "Bo cannot name Ada's chat to Test a question, read its Documents switch or change "
            "it (the same 404 each)",
            bo.post("/api/documents/test", json={"question": "code", "chatId": chat}).status_code
            == 404
            and bo.get(f"/api/chats/{chat}/documents").status_code == 404
            and bo.patch(f"/api/chats/{chat}", json={"settings": {"documents": "off"}}).status_code
            == 404,
        )
        bo_reply, bo_request, bo_reached = ask(bo, bo_chat, CODE_WORD_ASKED, word=BO_WORD)
        check(
            "Bo's chat with Documents on is sent only Bo's own file, at both hops, and the model "
            "answers from it",
            BO_WORD in json.dumps(bo_request["messages"])
            and BO_WORD in json.dumps(bo_reached)
            and CODE_TWO not in json.dumps(bo_request["messages"])
            and CODE_TWO not in json.dumps(bo_reached)
            and stem(BO_WORD) in bo_reply["content"].upper(),
            bo_reply["content"][:120],
        )
        same_words = "lighthouse inspection code word"
        bo_found = bo.post("/api/documents/test", json={"question": same_words}).json()
        ada_found = ada.post("/api/documents/test", json={"question": same_words}).json()
        check(
            "the same question finds Bo's file for Bo and Ada's file for Ada, and neither "
            "person's results hold a word of the other's",
            {p["name"] for p in bo_found["passages"]} == {"bo.md"}
            and {p["name"] for p in ada_found["passages"]} == {"survey.pdf"}
            and CODE_TWO not in json.dumps(bo_found)
            and BO_WORD not in json.dumps(ada_found),
            (bo_found["passages"][:1], ada_found["passages"][:1]),
        )
        check(
            "Bo's list is his own, and a stranger with no session is refused on every route",
            [d["name"] for d in bo.documents()["documents"]] == ["bo.md"]
            and all(
                httpx.request(
                    m, f"{WORKBENCH}{p}", headers={"Origin": WORKBENCH}, trust_env=False
                ).status_code
                == 401
                for m, p in (
                    ("GET", "/api/documents"),
                    ("GET", f"/api/documents/{mine['id']}"),
                    ("GET", "/api/documents/prefs"),
                    ("GET", f"/api/chats/{chat}/documents"),
                )
            ),
        )

        def documents_switch(person: Person) -> Any:
            made_now = person.new_chat(LOCAL)
            return person.get(f"/api/chats/{made_now}").json()["chat"]["settings"]["documents"]

        before_prefs = documents_switch(ada)
        said = ada.put("/api/documents/prefs", json={"newChats": True})
        ada_after, bo_after = documents_switch(ada), documents_switch(bo)
        ada.put("/api/documents/prefs", json={"newChats": None})
        check(
            "what a person says about new chats is theirs alone: off until they turn it on, on "
            "for their next chat, and Bo's chats are not touched",
            before_prefs is None
            and said.status_code == 200
            and said.json()["newChats"] is True
            and said.json()["effective"]["newChats"] is True
            and ada_after == "on"
            and bo_after is None
            and documents_switch(ada) is None
            and ada.get("/api/documents/prefs").json()["newChats"] is False,
            (before_prefs, said.text[:200], ada_after, bo_after),
        )
        check(
            "the owner cannot read Ada's documents, or her chats, while the owner's switch is off",
            owner.get(f"/api/documents/{target['id']}").status_code == 404
            and owner.get(f"/api/chats/{chat}").status_code == 404,
        )
        settings(ownerReadsChats=True)
        try:
            seen_chat = owner.get(f"/api/chats/{chat}")
            seen_messages = seen_chat.json()["messages"] if seen_chat.status_code == 200 else []
            hidden = [s for m in seen_messages for s in m["sources"] if s.get("passage")]
            check(
                "where the owner may read her chats, each passage of her documents is 'a "
                "document': no name, no words, no heading, no page",
                seen_chat.status_code == 200
                and bool(hidden)
                and all(s["name"] == "a document" and s["redacted"] is True for s in hidden)
                and not any(
                    {"snippet", "heading", "page_start", "doc_id", "chunk_ids"} & set(s)
                    for s in hidden
                )
                and "survey.pdf" not in json.dumps([m["sources"] for m in seen_messages])
                and "<documents-" not in seen_chat.text
                and all(
                    (m.get("retrieval") or {}).get("searched") is None
                    and (m.get("retrieval") or {}).get("cited") in (None, [])
                    for m in seen_messages
                ),
                hidden[:1],
            )
            check(
                "the owner still reads none of her documents: 404 on every route, an empty list "
                "of their own, and only the switch of her chat",
                owner.get(f"/api/documents/{target['id']}").status_code == 404
                and owner.get(f"/api/documents/{target['id']}/text").status_code == 404
                and owner.get(f"/api/documents/{target['id']}/original").status_code == 404
                and owner.documents()["documents"] == []
                and owner.documents()["ownerReadsChats"] is True
                and owner.get(f"/api/chats/{chat}/documents").json()
                == {"chatId": chat, "readOnly": True, "switch": "on", "documentsOnly": False},
            )
        finally:
            settings(ownerReadsChats=None)

    except Exception as exc:
        traceback.print_exc()
        check(f"the run went to plan ({type(exc).__name__}: {exc})", False)
    finally:
        say("11. teardown by pid")
        if server is not None:
            server.should_exit = True
        for proc in procs:
            subprocess.run(
                ["taskkill", "/PID", str(proc.pid), "/T", "/F"], capture_output=True, check=False
            )
        time.sleep(2)
        still = [p for p in OWNED if listening(p) and p not in OURS]
        check("no owned process port is still listening", not still, still)
        print(f"\nlogs: {logs}\n{PASSES} passed, {FAILURES} failed", flush=True)
    return 1 if FAILURES else 0


if __name__ == "__main__":
    raise SystemExit(main())
