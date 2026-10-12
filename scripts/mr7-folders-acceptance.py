"""Folders (MR7), live: a real model answers from a folder a person was granted, and the folder is kept
in step with it.

The unit checks drive Workbench against a fake gateway and a folder on a temporary disk. What only a live
run shows is the whole path: the owner grants a real folder, a real Workbench process walks it through held
handles and reads each file in a child process of its own, a real gateway, a real inference-driver and a
real `llama-server` answer from a real small model that has never seen the code words in the files, and the
folder is edited, a file is deleted, a link points outside it, and the grant is withdrawn, while the folder
itself is never written to. Two recording proxies keep what each hop was sent. The people are real
sign-ins against a stand-in for Eugene's sign-in (OpenID Connect, as the unit checks use).

  0. isolated: every ambient EUGENE_PLEXUS_* variable dropped, owned ports free
  1. the stack: engine (CPU, Qwen3 0.6B, a window of 8,192), a driver that says it is local, a gateway,
     Workbench as its own account (so that folders can be granted); the owner, Ada and Bo sign in
  2. the owner grants a folder to Ada; she adds it: its state while it is read, then every file's state,
     counts and name (the folder's name and its path), and the limits stated
  3. a question whose answer is a made-up code word in one file: the engine is sent the passage, the real
     model gives the word, and the Sources name the file by its folder and path
  4. what is left out is named: hidden, temporary and other programs' folders (with the rule), a picture and
     a program (each a document with its reason), a link that points outside (never followed)
  5. edit that file: after Sync now the new text answers and the old is gone from answers and from every file
     of the person's folder
  6. delete a file: it is gone from the list and from answers, and its words from every file
  7. a larger folder is read, and Workbench is killed in the middle of it: each file is read once, none left
     "reading", and what was read before is not read again
  8. two people on every route: Bo gets the same 404 for Ada's folder as for none, cannot add Ada's grant, and
     the owner cannot add it either
  9. the folder was never written to: every byte, name and time is as the run left it, apart from the edits
 10. the owner withdraws the grant: before the request returns the folder's documents are out of the index,
     answers and every file, the folder says so in words, and the person can dismiss it
 11. teardown by pid; no owned port still listening

**Safe beside a live install**: every ambient EUGENE_PLEXUS_* variable is dropped, ports are 8620-8627,
teardown is by pid. Never `pkill -f eugene_plexus_`: this box is a worker node.

Run with the Workbench interpreter (it has httpx, fastapi, uvicorn, joserfc, pypdf, python-docx):
  workbench/.venv/Scripts/python.exe workbench-mr/scripts/mr7-folders-acceptance.py
"""

from __future__ import annotations

import base64
import contextlib
import hashlib
import http.client
import json
import os
import re
import secrets
import socket
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import traceback
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any, cast
from urllib.parse import urlencode, urlsplit

import httpx
import uvicorn
from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse, RedirectResponse
from joserfc import jwt
from joserfc.jwk import RSAKey

ROOT = Path(os.environ.get("EP_ROOT", "D:/py/eugene-plexus"))
HERE = Path(__file__).resolve().parent.parent
#: The Workbench tree under test: the one this script sits in, else the main one.
WB = Path(
    os.environ.get(
        "EP_WORKBENCH",
        str(HERE if (HERE / "src" / "eugene_plexus_workbench").is_dir() else ROOT / "workbench"),
    )
)
WB_PY = Path(
    os.environ.get("EP_WORKBENCH_PY", str(ROOT / "workbench" / ".venv" / "Scripts" / "python.exe"))
)
GW_PY = ROOT / "gateway" / ".venv" / "Scripts" / "python.exe"
DRV_PY = ROOT / "inference-driver" / ".venv" / "Scripts" / "python.exe"
TEMP = Path(os.environ.get("TEMP", tempfile.gettempdir()))
LLAMA = Path(
    os.environ.get(
        "EP_LLAMA_SERVER", "C:/Users/troyc/.eugene-plexus/engines/llama_cpp/b10948/llama-server.exe"
    )
)
MODEL_FILE = Path(
    os.environ.get("EP_MODEL_FILE", str(TEMP / "ep-s9-model" / "Qwen3-0.6B-Q4_K_M.gguf"))
)

sys.path.insert(0, str(HERE))

(
    GW_PORT,
    AGENT_PORT,
    DRV_LOCAL,
    ENGINE,
    ENGINE_PROXY,
    WB_PORT,
    EUGENE_PORT,
    GW_PROXY,
) = range(8620, 8628)
OWNED = [GW_PORT, AGENT_PORT, DRV_LOCAL, ENGINE, ENGINE_PROXY, WB_PORT, EUGENE_PORT, GW_PROXY]
#: Ports a thread of this script serves; the rest belong to child processes.
OURS = (AGENT_PORT, ENGINE_PROXY, EUGENE_PORT, GW_PROXY)
WORKBENCH = f"http://127.0.0.1:{WB_PORT}"
LOCAL = "mr7-local"
FOLDER_NAME = "Work notes"

CLIENT_ID, CLIENT_SECRET, ADMIN_TOKEN = "mr7-client", "mr7-client-secret", "mr7-admin-token"
#: Made up, so only a model that was shown one can say it. Each file has its own words, so a question has
#: one right answer and the files do not compete for it.
CODE = "HERON-1180"  # budget/harbour-plan.md, as first written
CODE_TWO = "OSPREY-2255"  # the same file, once it is edited
PANTRY = "QUOKKA-7359"  # kitchen/pantry.md, until it is deleted
TIDE = "TERN-4417"  # tides.docx
OUTSIDE = "PLOVER-3081"  # a file outside the folder, behind a link inside it
CODE_WORD_ASKED = (
    "What is the code word for the lighthouse inspection? Reply with only the code word. /no_think"
)
PANTRY_ASKED = "What is the pantry padlock combination? Reply with only the combination. /no_think"
OUTSIDE_ASKED = (
    "What is the code word of the plover survey? Reply with only the code word. /no_think"
)

FAILURES = 0
PASSES = 0


def say(text: str) -> None:
    print(f"\n== {text}", flush=True)


def check(claim: str, passed: bool, detail: object = "") -> bool:
    global FAILURES, PASSES
    if passed:
        PASSES += 1
    else:
        FAILURES += 1
    shown = f"  ({str(detail)[:400]})" if detail != "" and not passed else ""
    print(f"  {'PASS' if passed else 'FAIL'}  {claim}{shown}", flush=True)
    return passed


def note(text: str) -> None:
    print(f"  NOTE  {text}", flush=True)


def listening(port: int) -> bool:
    with socket.socket() as s:
        s.settimeout(0.3)
        return s.connect_ex(("127.0.0.1", port)) == 0


def stem(word: str) -> str:
    return word.split("-")[0]


def harbour_plan(word: str) -> bytes:
    return (
        "# Harbour plan\n\nThe plan covers the quay, the lamps and the keeper's duties.\n\n"
        f"## Lighthouse inspection\n\nThe code word for the lighthouse inspection is {word}, "
        "told only to the inspector on the morning of the visit.\n\n"
        "## Berthing\n\nBerthing fees are paid at the quayside office each Friday.\n"
    ).encode()


def pantry_notes(word: str) -> bytes:
    return (
        "# Keeper's notes\n\nThe keeper writes about the garden and the kettle.\n\n"
        f"## The pantry\n\nThe pantry padlock combination is {word}, known to nobody but the old "
        "keeper.\n"
    ).encode()


def tides_docx(word: str) -> bytes:
    import io

    import docx as word_files

    document = word_files.Document()
    document.add_heading("Tide tables", level=1)
    document.add_paragraph("The tables follow the gauge at the harbour mouth.")
    document.add_heading("Dawn readings", level=2)
    document.add_paragraph(f"The tide gauge reading at dawn on the spring tide is {word}.")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def snapshot(folder: Path) -> dict[str, tuple[bytes | None, int]]:
    """Every file and directory below a folder, its bytes (None for a directory) and modified time. A link
    is one entry and is never followed: what it points at is not below the folder."""
    seen: dict[str, tuple[bytes | None, int]] = {}

    def walk(directory: Path) -> None:
        with os.scandir(directory) as entries:
            for entry in sorted(entries, key=lambda e: e.name):
                path = Path(entry.path)
                name = path.relative_to(folder).as_posix()
                if entry.is_symlink() or entry.is_junction():
                    seen[name] = (b"<link>", 0)
                elif entry.is_dir(follow_symlinks=False):
                    seen[name] = (None, entry.stat(follow_symlinks=False).st_mtime_ns)
                    walk(path)
                else:
                    seen[name] = (path.read_bytes(), entry.stat(follow_symlinks=False).st_mtime_ns)

    walk(folder)
    return seen


# --------------------------------------------------------------------------- #
# What each hop was sent
# --------------------------------------------------------------------------- #

#: Every chat request Workbench sent the gateway, and every one the engine received.
WB_SENT: list[dict[str, Any]] = []
ENGINE_BODIES: list[dict[str, Any]] = []
LOCK = threading.Lock()


class Quiet(ThreadingHTTPServer):
    daemon_threads = True

    def handle_error(self, request: Any, client_address: Any) -> None:
        # A probe that hangs up early is not news.
        pass


def make_proxy(upstream: int, sink: list[dict[str, Any]]) -> type[BaseHTTPRequestHandler]:
    """A pass-through to `127.0.0.1:upstream` that keeps the body of every chat request in `sink`:
    exactly what the next hop was sent."""

    class Proxy(BaseHTTPRequestHandler):
        def log_message(self, *args: Any) -> None:
            pass

        def _forward(self) -> None:
            length = int(self.headers.get("content-length") or 0)
            body = self.rfile.read(length) if length else b""
            if self.command == "POST" and self.path.split("?")[0] == "/v1/chat/completions":
                with LOCK:
                    sink.append(json.loads(body))
            hop = http.client.HTTPConnection("127.0.0.1", upstream, timeout=900)
            wanted = ("content-type", "accept", "authorization", "x-eugene-repetition-mode")
            headers = {k: v for k, v in self.headers.items() if k.lower() in wanted}
            hop.request(self.command, self.path, body=body or None, headers=headers)
            answer = hop.getresponse()
            self.send_response(answer.status)
            for key in ("Content-Type", "Cache-Control"):
                if value := answer.getheader(key):
                    self.send_header(key, value)
            self.end_headers()  # HTTP/1.0: the body ends when the connection does
            while chunk := answer.read(4096):
                self.wfile.write(chunk)
                self.wfile.flush()
            hop.close()

        do_GET = _forward
        do_POST = _forward
        do_PATCH = _forward
        do_DELETE = _forward

    return Proxy


class Agent(BaseHTTPRequestHandler):
    """Topology only: the gateway asks its agent which drivers there are."""

    def log_message(self, *args: Any) -> None:
        pass

    def do_GET(self) -> None:
        path = self.path.split("?")[0]
        if path == "/v1/components":
            body: Any = {
                "components": [
                    {
                        "name": "mr7-local-driver",
                        "kind": "inference-driver",
                        "url": f"http://127.0.0.1:{DRV_LOCAL}/",
                        "status": "running",
                    }
                ]
            }
        elif path == "/v1/runtimes":
            body = {"runtimes": []}
        elif path == "/v1/node":
            body = {"name": "mr7-node", "enrolled": False}
        else:
            body = {"status": "ok"}
        raw = json.dumps(body).encode()
        self.send_response(200)
        self.send_header("content-type", "application/json")
        self.send_header("content-length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)


def serve(handler: type[BaseHTTPRequestHandler], port: int) -> None:
    Quiet(("127.0.0.1", port), handler).serve_forever()


# --------------------------------------------------------------------------- #
# A stand-in for Eugene's sign-in (OpenID Connect, as the unit checks use)
# --------------------------------------------------------------------------- #

PEOPLE = {
    "operator": {"name": "Owner", "role": "operator", "username": "operator"},
    "p-ada": {"name": "Ada", "role": "member", "username": "ada"},
    "p-bo": {"name": "Bo", "role": "member", "username": "bo"},
}


def eugene_app() -> Any:
    key = RSAKey.generate_key(2048, parameters={"kid": "k1"})
    issuer = f"http://127.0.0.1:{EUGENE_PORT}/oidc"
    codes: dict[str, dict[str, Any]] = {}
    app = FastAPI()

    def id_token(sub: str, nonce: str | None) -> str:
        now = int(time.time())
        person = PEOPLE[sub]
        claims: dict[str, Any] = {
            "iss": issuer,
            "aud": CLIENT_ID,
            "sub": sub,
            "iat": now,
            "exp": now + 600,
            "name": person["name"],
            "preferred_username": person["username"],
            "eugene_role": person["role"],
            "auth_time": now,
        }
        if nonce is not None:
            claims["nonce"] = nonce
        return jwt.encode({"alg": "RS256", "typ": "JWT", "kid": "k1"}, claims, key)

    @app.get("/oidc/.well-known/openid-configuration")
    async def discovery() -> Any:
        return {
            "issuer": issuer,
            "authorization_endpoint": f"{issuer}/authorize",
            "token_endpoint": f"{issuer}/token",
            "jwks_uri": f"{issuer}/jwks",
            "revocation_endpoint": f"{issuer}/revoke",
        }

    @app.get("/oidc/jwks")
    async def jwks() -> Any:
        return {"keys": [key.as_dict(private=False)]}

    @app.get("/oidc/authorize")
    async def authorize(request: Request) -> Any:
        q = request.query_params
        code = secrets.token_urlsafe(16)
        codes[code] = {
            "sub": q["login_as"],
            "nonce": q["nonce"],
            "challenge": q["code_challenge"],
            "redirect_uri": q["redirect_uri"],
        }
        target = (
            q["redirect_uri"] + "?" + urlencode({"code": code, "state": q["state"], "iss": issuer})
        )
        return RedirectResponse(target, status_code=302)

    @app.post("/oidc/token")
    async def token(request: Request) -> Any:
        auth = request.headers.get("authorization", "")
        expected = base64.b64encode(f"{CLIENT_ID}:{CLIENT_SECRET}".encode()).decode()
        if auth != f"Basic {expected}":
            return JSONResponse({"error": "invalid_client"}, status_code=401)
        form = await request.form()
        if form["grant_type"] == "authorization_code":
            found = codes.pop(str(form["code"]), None)
            verifier = str(form.get("code_verifier", ""))
            challenge = (
                base64.urlsafe_b64encode(hashlib.sha256(verifier.encode()).digest())
                .rstrip(b"=")
                .decode()
            )
            if found is None or found["challenge"] != challenge:
                return JSONResponse({"error": "invalid_grant"}, status_code=400)
            return {
                "access_token": "at",
                "token_type": "Bearer",
                "expires_in": 3600,
                "id_token": id_token(found["sub"], found["nonce"]),
                "refresh_token": secrets.token_urlsafe(16),
            }
        return JSONResponse({"error": "unsupported_grant_type"}, status_code=400)

    @app.post("/oidc/revoke")
    async def revoke() -> Any:
        return {}

    return app


# --------------------------------------------------------------------------- #
# A person's browser
# --------------------------------------------------------------------------- #


class Person:
    """One person's browser: a cookie jar, the secret the page keeps, an origin."""

    def __init__(self, sub: str) -> None:
        self.sub = sub
        self.http = httpx.Client(base_url=WORKBENCH, timeout=240, trust_env=False)
        self.secret: str | None = None

    def sign_in(self) -> None:
        start = self.http.get("/signin")
        assert start.status_code == 302, start.text
        at_eugene = httpx.get(start.headers["location"] + "&login_as=" + self.sub, trust_env=False)
        assert at_eugene.status_code == 302, at_eugene.text
        back = urlsplit(at_eugene.headers["location"])
        done = self.http.get(back.path + "?" + back.query)
        location = done.headers.get("location", "")
        assert "#signin=" in location, f"sign-in failed: {location}"
        self.secret = location.split("#signin=", 1)[1]

    def call(self, method: str, path: str, **kw: Any) -> Any:
        headers = {"Origin": WORKBENCH, "X-Workbench-Secret": self.secret or ""}
        return self.http.request(method, path, headers=headers, **kw)

    def get(self, path: str, **kw: Any) -> Any:
        return self.call("GET", path, **kw)

    def post(self, path: str, **kw: Any) -> Any:
        return self.call("POST", path, **kw)

    def put(self, path: str, **kw: Any) -> Any:
        return self.call("PUT", path, **kw)

    def patch(self, path: str, **kw: Any) -> Any:
        return self.call("PATCH", path, **kw)

    def delete(self, path: str, **kw: Any) -> Any:
        return self.call("DELETE", path, **kw)

    def new_chat(self, model: str, **settings: Any) -> str:
        chat = self.post("/api/chats", json={"model": model}).json()["id"]
        if settings:
            changed = self.patch(f"/api/chats/{chat}", json={"settings": settings})
            assert changed.status_code == 200, changed.text
        return str(chat)

    # --- folders ---------------------------------------------------------

    def folder(self, grant: str) -> dict[str, Any]:
        found = self.get(f"/api/documents/folders/{grant}")
        assert found.status_code == 200, found.text
        return dict(found.json())

    def folders(self) -> dict[str, Any]:
        found = self.get("/api/documents/folders")
        assert found.status_code == 200, found.text
        return dict(found.json())

    def files(self, grant: str, **query: Any) -> dict[str, Any]:
        found = self.get(f"/api/documents/folders/{grant}/files", params={"limit": 500, **query})
        assert found.status_code == 200, found.text
        return dict(found.json())

    def by_path(self, grant: str) -> dict[str, dict[str, Any]]:
        return {f["path"]: f for f in self.files(grant)["files"]}

    def wait_folder(
        self, grant: str, *, after: float = 0.0, seconds: float = 300, until: Any = None
    ) -> dict[str, Any]:
        """The folder once a check that ended after `after` is done and every file has ended (any end),
        or once `until(folder)` says."""
        end = time.perf_counter() + seconds
        seen: dict[str, Any] = {}
        while time.perf_counter() < end:
            seen = self.folder(grant)
            by = seen["counts"]["byState"]
            reading = sum(by[s] for s in ("queued", "reading", "searchable", "embedding"))
            if until is not None:
                if until(seen):
                    return seen
            elif (
                seen["state"] in ("synced", "problem", "unavailable", "withdrawn")
                and (seen["checkedAt"] or 0.0) > after
                and seen["syncing"] is None
                and reading == 0
            ):
                return seen
            time.sleep(0.05)
        raise AssertionError(f"the folder was not read in time: {seen['state']} {seen['headline']}")

    def sync_now(self, grant: str, **body: Any) -> dict[str, Any]:
        before = float(self.folder(grant)["checkedAt"] or 0.0)
        made = self.post(f"/api/documents/folders/{grant}/sync", json=body)
        assert made.status_code == 202, made.text
        return self.wait_folder(grant, after=before)

    # --- documents --------------------------------------------------------

    def documents(self) -> dict[str, Any]:
        found = self.get("/api/documents")
        assert found.status_code == 200, found.text
        return dict(found.json())

    # --- chats ------------------------------------------------------------

    def wait_answer(self, chat: str, seconds: float = 300) -> dict[str, Any]:
        end = time.perf_counter() + seconds
        while time.perf_counter() < end:
            messages = self.get(f"/api/chats/{chat}").json()["messages"]
            if messages and messages[-1]["status"] != "running":
                return dict(messages[-1])
            time.sleep(0.3)
        raise AssertionError("the answer did not finish")

    def say(self, chat: str, content: str, **body: Any) -> Any:
        return self.post(f"/api/chats/{chat}/messages", json={"content": content, **body})


def newest(sink: list[dict[str, Any]], count: int, what: str) -> dict[str, Any]:
    """The newest request in `sink`, once it holds more than `count`."""
    end = time.time() + 60
    while time.time() < end:
        with LOCK:
            if len(sink) > count:
                return sink[-1]
        time.sleep(0.2)
    raise AssertionError(f"{what} was sent no request")


def text_of(message: dict[str, Any]) -> str:
    content = message["content"]
    return content if isinstance(content, str) else str(content[0]["text"])


BLOCK = re.compile(r"<documents-([0-9a-f]{8})>\n(.*?)\n</documents-\1>", re.DOTALL)
PASSAGE = re.compile(
    r'<document-([0-9a-f]{8}) n="(\d+)" name="([^"]*)"((?: [a-z]+="[^"]*")*)>'
    r"\n(.*?)\n</document-\1>",
    re.DOTALL,
)


def passages_in(message: dict[str, Any]) -> list[tuple[str, ...]]:
    return [m.groups() for m in PASSAGE.finditer(text_of(message))]


def person_key(sub: str) -> str:
    return hashlib.sha256(sub.encode("utf-8")).hexdigest()[:32]


# --------------------------------------------------------------------------- #
# The run
# --------------------------------------------------------------------------- #


def main() -> int:
    cast(Any, sys.stdout).reconfigure(encoding="utf-8", errors="replace")
    procs: list[subprocess.Popen[bytes]] = []
    work = Path(tempfile.mkdtemp(prefix="ep-mr7-"))
    # The granted folder, and a folder outside it that a link inside it points at: neither is under
    # Workbench's own files, which cannot be granted.
    folder = Path(tempfile.mkdtemp(prefix="ep-mr7-notes-"))
    outside = Path(tempfile.mkdtemp(prefix="ep-mr7-outside-"))
    logs = work / "logs"
    logs.mkdir()
    data = work / "workbench-data"
    wb_log_count = [0]

    def start(argv: list[str], log: str, env: dict[str, str]) -> subprocess.Popen[bytes]:
        # The child owns the handle for as long as it runs; teardown kills it.
        handle = open(logs / log, "w", encoding="utf-8")  # noqa: SIM115
        proc = subprocess.Popen(
            argv, stdin=subprocess.DEVNULL, stdout=handle, stderr=subprocess.STDOUT, env=env
        )
        procs.append(proc)
        return proc

    def wait_for(port: int, seconds: float = 120) -> bool:
        end = time.time() + seconds
        while time.time() < end:
            if listening(port):
                return True
            time.sleep(0.5)
        return False

    def kill(proc: subprocess.Popen[bytes]) -> None:
        subprocess.run(
            ["taskkill", "/PID", str(proc.pid), "/T", "/F"], capture_output=True, check=False
        )

    def counts() -> tuple[int, int]:
        with LOCK:
            return len(WB_SENT), len(ENGINE_BODIES)

    def index_path(sub: str) -> Path:
        return data / "knowledge" / person_key(sub) / "index.sqlite3"

    def knowledge_files(sub: str) -> list[Path]:
        root = data / "knowledge" / person_key(sub)
        return [p for p in root.rglob("*") if p.is_file()] if root.exists() else []

    def holds(sub: str, word: str) -> list[str]:
        """Every file of a person's own folder that holds this word, as it was written or as the full-text
        index keeps it: the index, its log, an original, the text read from it, what is said of it."""
        needles = (word.lower().encode(), stem(word).lower().encode())
        return [
            p.name
            for p in knowledge_files(sub)
            if not p.name.endswith("-shm") and any(n in p.read_bytes().lower() for n in needles)
        ]

    def holds_whole(sub: str, word: str) -> list[str]:
        """The same, for a word whose first half is also a piece of other words (`tern` is in
        `external`): only the whole code word counts."""
        needle = word.lower().encode()
        return [
            p.name
            for p in knowledge_files(sub)
            if not p.name.endswith("-shm") and needle in p.read_bytes().lower()
        ]

    def age(path: Path) -> None:
        """Every file below `path` that was written lately was last modified an hour ago, so none is
        too new to trust. A file already aged is left as it is: its modified time is what it was."""
        now = time.time()
        for found in path.rglob("*"):
            if found.is_file() and found.stat().st_mtime > now - 1800:
                os.utime(found, (now - 3600, now - 3600))

    def write(name: str, content: bytes) -> Path:
        target = folder / name
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(content)
        return target

    server: Any = None
    wb: list[subprocess.Popen[bytes]] = []
    try:
        say("0. isolate this run")
        base = {k: v for k, v in os.environ.items() if not k.startswith("EUGENE_PLEXUS_")}
        for proxy in ("HTTP_PROXY", "HTTPS_PROXY", "ALL_PROXY", "NO_PROXY"):
            base.pop(proxy, None)
            base.pop(proxy.lower(), None)
        check("no ambient EUGENE_PLEXUS_* variable reaches a child", True)
        taken = [p for p in OWNED if listening(p)]
        if taken:
            check(f"ports {taken} are free", False)
            return 1
        for path in (
            WB_PY,
            GW_PY,
            DRV_PY,
            LLAMA,
            MODEL_FILE,
            WB / "src" / "eugene_plexus_workbench",
        ):
            if not path.exists():
                check(f"{path} exists", False)
                return 1
        check(
            f"ports {OWNED[0]}-{OWNED[-1]} free; Workbench from {WB}; llama-server "
            f"{LLAMA.parent.name}; {MODEL_FILE.name}",
            True,
        )

        say("1. the stack: engine, driver, gateway, Workbench, and three people")
        threading.Thread(target=serve, args=(Agent, AGENT_PORT), daemon=True).start()
        threading.Thread(
            target=serve, args=(make_proxy(ENGINE, ENGINE_BODIES), ENGINE_PROXY), daemon=True
        ).start()
        threading.Thread(
            target=serve, args=(make_proxy(GW_PORT, WB_SENT), GW_PROXY), daemon=True
        ).start()
        config = uvicorn.Config(eugene_app(), host="127.0.0.1", port=EUGENE_PORT, log_level="error")
        server = uvicorn.Server(config)
        threading.Thread(target=server.run, daemon=True).start()
        # CPU only, one slot (several would divide the window), a window of 8,192.
        engine = [str(LLAMA), "-m", str(MODEL_FILE), "--host", "127.0.0.1", "--port", str(ENGINE)]
        engine += ["-ngl", "0", "-t", "4", "-c", "8192", "-np", "1", "-b", "512", "-ub", "512"]
        start(
            engine,
            "llama.log",
            {**base, "CUDA_VISIBLE_DEVICES": "-1", "GGML_VK_VISIBLE_DEVICES": ""},
        )
        (work / "driver-local.yaml").write_text(
            f"provider: openai_compat_custom\nbaseUrl: http://127.0.0.1:{ENGINE_PROXY}\n"
            f"modelId: {LOCAL}\nbackendLocality: local\nrequestTimeoutSeconds: 600\n",
            encoding="utf-8",
            newline="\n",
        )
        start(
            [str(DRV_PY), "-m", "eugene_plexus_inference_driver"],
            "driver-local.log",
            {
                **base,
                "EUGENE_PLEXUS_DRIVER_CONFIG_FILE": str(work / "driver-local.yaml"),
                "EUGENE_PLEXUS_DRIVER_BIND_HOST": "127.0.0.1",
                "EUGENE_PLEXUS_DRIVER_BIND_PORT": str(DRV_LOCAL),
            },
        )
        (work / "gateway.yaml").write_text("routingRefreshSeconds: 3\n", encoding="utf-8")
        start(
            [str(GW_PY), "-m", "eugene_plexus_gateway"],
            "gateway.log",
            {
                **base,
                "EUGENE_PLEXUS_GATEWAY_CONFIG_FILE": str(work / "gateway.yaml"),
                "EUGENE_PLEXUS_GATEWAY_METRICS_FILE": str(work / "metrics.sqlite3"),
                "EUGENE_PLEXUS_GATEWAY_BIND_HOST": "127.0.0.1",
                "EUGENE_PLEXUS_GATEWAY_BIND_PORT": str(GW_PORT),
                "EUGENE_PLEXUS_GATEWAY_AGENT_URL": f"http://127.0.0.1:{AGENT_PORT}",
            },
        )
        for port in (ENGINE, DRV_LOCAL, GW_PORT, EUGENE_PORT):
            if not wait_for(port, 180):
                check(f"something listens on {port}; logs in {logs}", False)
                return 1
        ids: set[str] = set()
        known = False
        end = time.time() + 240
        while time.time() < end and not known:
            try:
                listed = httpx.get(
                    f"http://127.0.0.1:{GW_PORT}/v1/models", timeout=5, trust_env=False
                ).json()
                ids = {m["id"] for m in listed.get("data") or []}
                entry = next((m for m in listed.get("data") or [] if m["id"] == LOCAL), None)
                # The driver learns the engine's window once the model has loaded.
                known = (
                    entry is not None
                    and (entry.get("x_eugene_plexus") or {}).get("context_length") is not None
                )
            except (httpx.HTTPError, ValueError):
                pass
            time.sleep(1)
        if not check(f"the gateway routes {LOCAL} and knows its window", known, sorted(ids)):
            return 1

        static = work / "static"
        static.mkdir()
        (static / "index.html").write_text(
            "<!doctype html><title>Workbench</title>", encoding="utf-8"
        )
        (work / "app_key").write_text("mr7-app-key", encoding="utf-8")
        (work / "oidc_secret").write_text(CLIENT_SECRET, encoding="utf-8")

        def start_workbench() -> subprocess.Popen[bytes]:
            wb_log_count[0] += 1
            return start(
                [str(WB_PY), "-m", "eugene_plexus_workbench"],
                f"workbench-{wb_log_count[0]}.log",
                {
                    **base,
                    "PYTHONPATH": str(WB / "src"),
                    "EUGENE_PLEXUS_APP_BIND_HOST": "127.0.0.1",
                    "EUGENE_PLEXUS_APP_BIND_PORT": str(WB_PORT),
                    "EUGENE_PLEXUS_APP_DATA_DIR": str(data),
                    "EUGENE_PLEXUS_APP_KEY_FILE": str(work / "app_key"),
                    "EUGENE_PLEXUS_APP_ADMIN_TOKEN": ADMIN_TOKEN,
                    # Through the recording proxy: it is what Workbench sends the gateway.
                    "EUGENE_PLEXUS_APP_GATEWAY_URL": f"http://127.0.0.1:{GW_PROXY}",
                    "EUGENE_PLEXUS_APP_OIDC_ISSUER": f"http://127.0.0.1:{EUGENE_PORT}/oidc",
                    "EUGENE_PLEXUS_APP_OIDC_CLIENT_ID": CLIENT_ID,
                    "EUGENE_PLEXUS_APP_OIDC_SECRET_FILE": str(work / "oidc_secret"),
                    "EUGENE_PLEXUS_APP_STATIC_DIR": str(static),
                    # Folders are granted only to Workbench's own account (the agent says so when it
                    # starts Workbench as a service); a run on a developer's own account says it too.
                    "EUGENE_PLEXUS_APP_ACCOUNT_KIND": "windows_service",
                },
            )

        wb.append(start_workbench())
        if not wait_for(WB_PORT, 60):
            check(f"Workbench listens on {WB_PORT}; logs in {logs}", False)
            return 1
        ada, bo, owner = Person("p-ada"), Person("p-bo"), Person("operator")
        for person in (ada, bo, owner):
            person.sign_in()
        check("Ada, Bo and the owner sign in with Eugene", all(p.secret for p in (ada, bo, owner)))
        listing = {m["id"]: m for m in ada.get("/api/models").json()["models"]}
        window = listing[LOCAL]["contextLength"]
        note(f"{LOCAL}: window {window} tokens; locality {listing[LOCAL]['locality']}")
        if not check(
            "the model is a local one with a window of 8,192",
            window == 8192 and listing[LOCAL]["locality"] == "local",
            listing[LOCAL],
        ):
            return 1

        def settings(**values: Any) -> Any:
            return httpx.patch(
                f"{WORKBENCH}/v1/config",
                json=values,
                headers={"Authorization": f"Bearer {ADMIN_TOKEN}"},
                trust_env=False,
            )

        def reached_by(reply: dict[str, Any], before_engine: int) -> dict[str, Any]:
            """What the engine received for this answer; nothing when no model was asked."""
            if reply["status"] != "done" or reply.get("finish") == "documents_only":
                return {}
            return newest(ENGINE_BODIES, before_engine, "the engine")

        def ask(
            person: Person,
            chat: str,
            question: str,
            *,
            word: str | None = None,
            tries: int = 3,
            **body: Any,
        ) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
            """Ask, wait for the real model, and ask again (a small model is not perfect) until its
            answer holds `word`. The answer, what Workbench sent the gateway, and what reached the
            engine, for the last try."""
            before_sent, before_engine = counts()
            sent = person.say(chat, question, **body)
            assert sent.status_code == 201, sent.text
            reply = person.wait_answer(chat)
            request = newest(WB_SENT, before_sent, "the gateway")
            reached = reached_by(reply, before_engine)
            tried = 1
            while word is not None and stem(word) not in reply["content"].upper() and tried < tries:
                tried += 1
                before_sent, before_engine = counts()
                person.post(f"/api/chats/{chat}/retry")
                reply = person.wait_answer(chat)
                request = newest(WB_SENT, before_sent, "the gateway")
                reached = reached_by(reply, before_engine)
            return reply, request, reached

        # ------------------------------------------------------------------ #
        say("2. the owner grants a folder to Ada; she adds it")
        write("budget/harbour-plan.md", harbour_plan(CODE))
        write("kitchen/pantry.md", pantry_notes(PANTRY))
        write("tides.docx", tides_docx(TIDE))
        write("notes/readme.txt", b"Plain notes about the garden, the kettle and the weather.\n")
        age(folder)
        before_everything = snapshot(folder)
        made = owner.post(
            "/api/folders",
            json={"name": FOLDER_NAME, "path": str(folder), "subject": "p-ada", "writable": False},
        )
        check(
            "the owner grants the folder to Ada, read-only",
            made.status_code == 201 and made.json()["subject"] == "p-ada",
            made.text[:200],
        )
        grant = made.json()["id"]
        offered = ada.folders()
        check(
            "Ada is offered it, by its name and never its path, with every limit stated",
            offered["available"] is True
            and offered["grants"]
            == [{"id": grant, "name": FOLDER_NAME, "writable": False, "added": False}]
            and str(folder) not in json.dumps(offered)
            and offered["limits"]["maxFiles"] == 10_000
            and offered["limits"]["maxMB"] == 2048
            and offered["limits"]["maxDepth"] == 12
            and offered["limits"]["syncMinutes"] == 10
            and ".git" in offered["limits"]["rules"]["folders"],
            offered["limits"],
        )
        added = ada.post("/api/documents/folders", json={"grantId": grant})
        first = added.json()
        seen_states: list[str] = [first["state"]]
        end = time.perf_counter() + 300
        view = first
        while time.perf_counter() < end:
            view = ada.folder(grant)
            if view["state"] != seen_states[-1]:
                seen_states.append(view["state"])
            by = view["counts"]["byState"]
            reading = sum(by[s] for s in ("queued", "reading", "searchable", "embedding"))
            if view["state"] == "synced" and view["syncing"] is None and reading == 0:
                break
            time.sleep(0.01)
        note("the folder's states: " + " -> ".join(seen_states))
        got = ada.by_path(grant)
        check(
            "Ada adds it (201): it is looked at at once, goes through its states, and ends synced",
            added.status_code == 201
            and first["state"] in ("waiting", "syncing")
            and seen_states[-1] == "synced"
            and view["problem"] is None
            and view["checkedAt"] is not None,
            (added.text[:200], seen_states),
        )
        check(
            "each readable file is a document named by the folder and its path, ready, with its "
            "counts; the files are not in the list of what she added herself",
            {p: d["state"] for p, d in got.items() if d["state"] == "ready"}
            == dict.fromkeys(
                ["budget/harbour-plan.md", "kitchen/pantry.md", "tides.docx", "notes/readme.txt"],
                "ready",
            )
            and got["budget/harbour-plan.md"]["name"] == f"{FOLDER_NAME}/budget/harbour-plan.md"
            and got["tides.docx"]["label"] == "Word"
            and all(d["passages"] >= 1 and d["tokens"] for d in got.values())
            and ada.documents()["documents"] == [],
            got,
        )
        check(
            "the headline says when it was checked, how often, and what it holds; the caps are "
            "shown with how near it is",
            view["headline"].startswith("Checked ")
            and "every 10 minutes" in view["headline"]
            and "4 files: 4 ready" in view["headline"]
            and view["caps"]["files"] == {"limit": 10_000, "found": 4, "reached": False}
            and 0 < view["caps"]["bytes"]["used"] < view["caps"]["bytes"]["limit"]
            and view["nextAt"] is not None,
            view["headline"],
        )
        search = ada.documents()["search"]
        check(
            "no search model is running, and the screen says documents are searched by words only",
            search["mode"] == "words" and search["headline"].startswith("Searching by words only"),
            search,
        )

        # ------------------------------------------------------------------ #
        say("3. a question answered from the folder: the Sources name the file")
        chat = ada.new_chat(LOCAL, temperature=0, maxTokens=300, documents="on")
        reply, request, reached = ask(ada, chat, CODE_WORD_ASKED, word=CODE)
        latest = request["messages"][-1]
        found_passages = passages_in(latest)
        check(
            "Workbench's request is ONE user message: the passages block first, the question last",
            [m["role"] for m in request["messages"]] == ["user"]
            and BLOCK.search(text_of(latest)) is not None
            and text_of(latest).startswith("<documents-")
            and text_of(latest).endswith("\n\n" + CODE_WORD_ASKED),
            [m["role"] for m in request["messages"]],
        )
        check(
            "the first passage is the folder's file, named by its folder and path, with the code word",
            bool(found_passages)
            and found_passages[0][2] == f"{FOLDER_NAME}/budget/harbour-plan.md"
            and CODE in found_passages[0][4],
            found_passages[:1],
        )
        check(
            "and the same passage reached the real engine, through the gateway and the driver",
            bool(reached) and CODE in json.dumps(reached["messages"]),
        )
        check(
            "the real model answers with the code word it was given",
            stem(CODE) in reply["content"].upper(),
            reply["content"][:150],
        )
        note(f"answer: {reply['content'][:100]!r}")
        passages = [s for s in reply["sources"] if s.get("passage")]
        check(
            "the Sources entry names the file, its heading and how it was found",
            bool(passages)
            and passages[0]["kind"] == "document"
            and passages[0]["name"] == f"{FOLDER_NAME}/budget/harbour-plan.md"
            and "Lighthouse inspection" in (passages[0]["heading"] or "")
            and passages[0]["via"] == "words"
            and CODE in passages[0]["snippet"],
            passages[:1],
        )
        shown = ada.get(f"/api/documents/{got['budget/harbour-plan.md']['id']}/text")
        original = ada.get(f"/api/documents/{got['budget/harbour-plan.md']['id']}/original")
        check(
            "Show indexed text and Open original work for a folder's file, and the download is the "
            "file's own name",
            shown.status_code == 200
            and CODE in shown.text
            and original.status_code == 200
            and original.content == harbour_plan(CODE)
            and 'filename="harbour-plan.md"' in original.headers["content-disposition"],
        )

        # ------------------------------------------------------------------ #
        say("4. what is left out is named, with its reason")
        write(".hidden-notes.md", b"# Hidden\n\nA hidden file nobody asked to read.\n")
        write("node_modules/pkg/readme.md", b"# Package\n\nAnother program's folder.\n")
        write("scratch.tmp", b"scratch")
        write("diagram.png", b"\x89PNG\r\n\x1a\n" + bytes(64))
        write("program.exe", b"MZ\x90\x00" + bytes(300))
        (outside / "secret.md").write_bytes(
            f"# Plover survey\n\nThe code word of the plover survey is {OUTSIDE}.\n".encode()
        )
        subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(folder / "shortcut"), str(outside)],
            check=True,
            capture_output=True,
        )
        age(folder)
        age(outside)
        outside_before = snapshot(outside)
        done = ada.sync_now(grant)
        shown_folder = ada.folder(grant)
        left = {x["path"]: x for x in shown_folder["left"]}
        got = ada.by_path(grant)
        check(
            "the hidden file, the temporary file and the other program's folder are named with the "
            "rule that left them out, and none is a document",
            left[".hidden-notes.md"]["reason"] == "Ignored: a hidden file."
            and left["scratch.tmp"]["code"] == "ignored_temporary"
            and left["node_modules"]["code"] == "ignored_default"
            and ".hidden-notes.md" not in got
            and "scratch.tmp" not in got
            and not any(p.startswith("node_modules") for p in got),
            left,
        )
        check(
            "a picture and a program each are a document of their own, skipped, with their reasons "
            "in words and a code the page can act on",
            got["diagram.png"]["state"] == "skipped"
            and got["diagram.png"]["code"] == "not_document"
            and got["diagram.png"]["message"].startswith(
                "Skipped: this file is a picture, not a document."
            )
            and got["program.exe"]["state"] == "skipped"
            and got["program.exe"]["code"] == "unsupported"
            and "binary file" in got["program.exe"]["message"]
            and done["counts"]["byCode"] == {"not_document": 1, "unsupported": 1}
            and ada.files(grant, state="skipped")["total"] == 2,
            (got["diagram.png"], got["program.exe"]),
        )
        reply_out, request_out, _ = ask(ada, chat, OUTSIDE_ASKED)
        found_out = ada.post("/api/documents/test", json={"question": OUTSIDE_ASKED}).json()
        check(
            "a link inside the folder that points outside it is named, not followed: the file it "
            "leads to is in no answer, no search and no file of the person's folder",
            left["shortcut"]["code"] == "link"
            and "never follows links" in left["shortcut"]["reason"]
            and done["last"]["links"] == 1
            and OUTSIDE not in json.dumps(request_out["messages"])
            and OUTSIDE not in reply_out["content"]
            and all(OUTSIDE not in p["text"] for p in found_out["passages"])
            and holds("p-ada", OUTSIDE) == []
            and snapshot(outside) == outside_before,
            (left.get("shortcut"), holds("p-ada", OUTSIDE)),
        )

        # ------------------------------------------------------------------ #
        say("5. edit the file: the new words answer, the old are gone from everything")
        old_id = got["budget/harbour-plan.md"]["id"]
        check(
            "before the edit the old code word is in the person's own files (so that the check after "
            "it means something)",
            holds("p-ada", CODE) != [],
        )
        write("budget/harbour-plan.md", harbour_plan(CODE_TWO))
        edit_chat = ada.new_chat(LOCAL, temperature=0, maxTokens=300, documents="on")
        done = ada.sync_now(grant)
        got = ada.by_path(grant)
        reply2, request2, reached2 = ask(ada, edit_chat, CODE_WORD_ASKED, word=CODE_TWO)
        check(
            "after Sync now the file is read again as a new document, and the old one is gone at once",
            got["budget/harbour-plan.md"]["id"] != old_id
            and got["budget/harbour-plan.md"]["state"] == "ready"
            and done["last"]["changed"] == 1
            and ada.get(f"/api/documents/{old_id}").status_code == 404,
            (done["last"], got["budget/harbour-plan.md"]["state"]),
        )
        check(
            "the new text is what the model is sent and what it answers with, and the old code word "
            "is nowhere in the request at either hop",
            CODE_TWO in json.dumps(request2["messages"])
            and CODE_TWO in json.dumps(reached2)
            and stem(CODE_TWO) in reply2["content"].upper()
            and CODE not in json.dumps(request2["messages"])
            and CODE not in json.dumps(reached2)
            and CODE not in reply2["content"],
            reply2["content"][:120],
        )
        old_found = ada.post(
            "/api/documents/test", json={"question": CODE_WORD_ASKED.replace(" /no_think", "")}
        ).json()
        check(
            "Test a question finds the new passage and no passage holds the old word",
            old_found["status"] == "found"
            and all(CODE not in p["text"] for p in old_found["passages"])
            and any(CODE_TWO in p["text"] for p in old_found["passages"]),
        )
        check(
            "and the old word is in no file of the person's folder: not the index, nor its log, nor "
            "an original, nor a text",
            holds("p-ada", CODE) == [],
            holds("p-ada", CODE),
        )

        # ------------------------------------------------------------------ #
        say("6. delete a file: it is gone from the list, from answers and from every file")
        pantry_id = got["kitchen/pantry.md"]["id"]
        pantry_chat = ada.new_chat(LOCAL, temperature=0, maxTokens=300, documents="on")
        before_reply, before_request, _ = ask(ada, pantry_chat, PANTRY_ASKED, word=PANTRY)
        check(
            "before it is deleted, the model answers the pantry question from the file",
            PANTRY in json.dumps(before_request["messages"])
            and stem(PANTRY) in before_reply["content"].upper(),
            before_reply["content"][:100],
        )
        (folder / "kitchen" / "pantry.md").unlink()
        done = ada.sync_now(grant)
        got = ada.by_path(grant)
        strict = ada.new_chat(
            LOCAL, temperature=0, maxTokens=300, documents="on", documentsOnly=True
        )
        sent_strict = ada.say(strict, PANTRY_ASKED)
        assert sent_strict.status_code == 201, sent_strict.text
        gone_reply = ada.wait_answer(strict)  # no model is asked: no request reaches the gateway
        check(
            "the next check removes it: it is not in the list, its id is a 404, and the check says so",
            "kitchen/pantry.md" not in got
            and ada.get(f"/api/documents/{pantry_id}").status_code == 404
            and done["last"]["removed"] == 1,
            done["last"],
        )
        check(
            "a chat that answers only from documents says nothing matched in Workbench's own words, "
            "and no model was asked",
            gone_reply.get("finish") == "documents_only"
            and gone_reply["retrieval"]["status"] == "nothing_matched"
            and PANTRY not in gone_reply["content"],
            gone_reply["content"][:160],
        )
        check(
            "its words are in no file of the person's folder",
            holds("p-ada", PANTRY) == [],
            holds("p-ada", PANTRY),
        )

        # ------------------------------------------------------------------ #
        say("7. a larger folder is read, and Workbench is killed in the middle of it")
        for n in range(40):
            write(
                f"big/note-{n:02}.md",
                f"# Note {n}\n\nNote {n} is about the weather, the tide and the gulls.\n".encode(),
            )
        age(folder)
        before_kill = {p: d["indexedAt"] for p, d in ada.by_path(grant).items() if d["indexedAt"]}
        started_at = time.perf_counter()
        request_sent = ada.post(f"/api/documents/folders/{grant}/sync", json={})
        mid = ada.wait_folder(
            grant,
            until=lambda f: (
                f["counts"]["byState"]["ready"] >= 8 + len(before_kill)
                and f["counts"]["files"] >= 20
            ),
            seconds=300,
        )
        ready_then = mid["counts"]["byState"]["ready"]
        for proc in wb:
            kill(proc)
        time.sleep(1)
        wb.clear()
        wb.append(start_workbench())
        ok = wait_for(WB_PORT, 60)
        for person in (ada, bo, owner):
            person.sign_in()
        end_state = ada.wait_folder(grant, seconds=600)
        elapsed = time.perf_counter() - started_at
        after = ada.by_path(grant)
        with contextlib.closing(sqlite3.connect(index_path("p-ada"))) as db:
            attempts = dict(db.execute("SELECT source, attempt FROM documents WHERE attempt > 1"))
            total_rows = db.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
        check(
            "Workbench was stopped part-way (some were read, some not) and started again",
            request_sent.status_code == 202 and ok and 0 < ready_then < 40 + len(before_kill),
            (ready_then, ok),
        )
        big = {p: d for p, d in after.items() if p.startswith("big/")}
        check(
            "after the restart every file of the larger folder is read, once, and none is left saying "
            "'reading'",
            len(big) == 40
            and {d["state"] for d in big.values()} == {"ready"}
            and end_state["state"] == "synced"
            and end_state["counts"]["byState"]["queued"] == 0
            and end_state["counts"]["byState"]["reading"] == 0
            and total_rows == len(after),
            (len(big), end_state["counts"]["byState"], total_rows),
        )
        check(
            "and what was read before was not read again: the files that were ready keep the time "
            "they were read",
            all(after[p]["indexedAt"] == t for p, t in before_kill.items() if p in after),
            [p for p, t in before_kill.items() if p in after and after[p]["indexedAt"] != t][:3],
        )
        note(
            f"40 files and the restart took {elapsed:.0f} s; {len(after)} documents; "
            f"documents read twice: {len(attempts)}"
        )

        # ------------------------------------------------------------------ #
        say("8. two people on every route, and the owner")
        missing = "doesnotexist1"
        same = True
        detail: list[Any] = []
        for method, route, body in (
            ("get", "/api/documents/folders/{id}", {}),
            ("patch", "/api/documents/folders/{id}", {"json": {"includeHidden": True}}),
            ("post", "/api/documents/folders/{id}/sync", {"json": {}}),
            ("get", "/api/documents/folders/{id}/files", {}),
            ("delete", "/api/documents/folders/{id}", {}),
        ):
            theirs = getattr(bo, method)(route.format(id=grant), **body)
            nothing = getattr(bo, method)(route.format(id=missing), **body)
            ok2 = (
                theirs.status_code == nothing.status_code == 404 and theirs.json() == nothing.json()
            )
            same = same and ok2
            detail.append((method, route, theirs.status_code, nothing.status_code))
        check(
            "Bo gets the same 404 for Ada's folder as for one that does not exist, on every route "
            "that takes a folder",
            same,
            detail,
        )
        refused_bo = bo.post("/api/documents/folders", json={"grantId": grant})
        refused_owner = owner.post("/api/documents/folders", json={"grantId": grant})
        none_bo = bo.post("/api/documents/folders", json={"grantId": missing})
        check(
            "Bo cannot add Ada's grant, and nor can the owner: the same 404 as a grant that does "
            "not exist",
            refused_bo.status_code == refused_owner.status_code == none_bo.status_code == 404
            and refused_bo.json() == refused_owner.json() == none_bo.json(),
            (refused_bo.text, refused_owner.text),
        )
        bo_chat = bo.new_chat(LOCAL, temperature=0, maxTokens=300, documents="on")
        _, bo_request, _ = ask(bo, bo_chat, CODE_WORD_ASKED)
        check(
            "Bo's chat with Documents on is sent none of Ada's folder, and Bo's list holds no folder",
            "<documents-" not in json.dumps(bo_request["messages"])
            and CODE_TWO not in json.dumps(bo_request["messages"])
            and bo.folders()["folders"] == []
            and bo.folders()["grants"] == []
            and bo.documents()["documents"] == [],
        )
        owner_view = owner.folders()
        check(
            "the owner is offered no grant that does not name the owner, and still sees every "
            "path in the folders list that has always shown them",
            owner_view["folders"] == []
            and owner_view["grants"] == []
            and owner.get("/api/folders").json()["grants"][0]["path"] == str(folder),
        )
        settings(ownerReadsChats=True)
        try:
            seen_chat = owner.get(f"/api/chats/{chat}")
            hidden = [
                s
                for m in (seen_chat.json()["messages"] if seen_chat.status_code == 200 else [])
                for s in m["sources"]
                if s.get("passage")
            ]
            check(
                "where the owner may read Ada's chats, each passage of her folder's files is 'a "
                "document': no name, no words, no path",
                seen_chat.status_code == 200
                and bool(hidden)
                and all(s["name"] == "a document" and s["redacted"] is True for s in hidden)
                and "harbour-plan" not in seen_chat.text
                and FOLDER_NAME not in seen_chat.text,
                hidden[:1],
            )
        finally:
            settings(ownerReadsChats=None)

        # ------------------------------------------------------------------ #
        say("9. the folder was never written to")
        after_everything = snapshot(folder)
        # What the run itself put there, and what it took away: nothing else is different.
        run_added = {
            ".hidden-notes.md",
            "node_modules",
            "node_modules/pkg",
            "node_modules/pkg/readme.md",
            "scratch.tmp",
            "diagram.png",
            "program.exe",
            "shortcut",
            "big",
            *(f"big/note-{n:02}.md" for n in range(40)),
        }
        run_changed = {"budget/harbour-plan.md", "kitchen", "kitchen/pantry.md"}
        unchanged = {p: v for p, v in before_everything.items() if p not in run_changed}
        different = [p for p, v in unchanged.items() if after_everything.get(p) != v]
        extra = set(after_everything) - set(before_everything) - run_added
        check(
            "every file the run did not itself change is exactly as it was, to the byte and the "
            "modified time, and nothing was created in the folder that the run did not make",
            not different and not extra,
            (different, sorted(extra)),
        )

        # ------------------------------------------------------------------ #
        say("10. the owner withdraws the grant")
        folder_before = snapshot(folder)
        withdrawn = owner.delete(f"/api/folders/{grant}")
        # No waiting: the documents were out of the index when that request returned.
        after_withdraw = ada.folder(grant)
        found_now = ada.post("/api/documents/test", json={"question": CODE_WORD_ASKED}).json()
        chip = ada.get(f"/api/chats/{edit_chat}/documents").json()
        check(
            "the owner's request returned, and the folder's documents are already out: no passage, no "
            "file, none counted",
            withdrawn.status_code == 204
            and found_now["status"] == "no_documents"
            and found_now["passages"] == []
            and chip["counts"]["library"]["documents"] == 0
            and ada.files(grant)["total"] == 0,
            (found_now["status"], chip["counts"]),
        )
        removed_n = after_withdraw["problem"]["removed"] if after_withdraw["problem"] else None
        check(
            "the folder stays in her list as withdrawn and says so in words, with how many documents "
            "were taken out",
            after_withdraw["state"] == "withdrawn"
            and after_withdraw["granted"] is False
            and removed_n == len(after)
            and after_withdraw["headline"]
            == f"The owner took away your access to this folder. {removed_n} documents from it "
            "were taken out of your search."
            and ada.folders()["grants"] == [],
            after_withdraw["headline"],
        )
        # A new chat: the one that was asked before holds the answer it was given, in its own words.
        after_chat = ada.new_chat(LOCAL, temperature=0, maxTokens=300, documents="on")
        reply3, request3, _ = ask(ada, after_chat, CODE_WORD_ASKED)
        check(
            "a chat that asks a question the folder answered is sent nothing from it now, and no "
            "word of the folder is in any file of her folder",
            "<documents-" not in json.dumps(request3["messages"])
            and CODE_TWO not in json.dumps(request3["messages"])
            and stem(CODE_TWO) not in reply3["content"].upper()
            and holds("p-ada", CODE_TWO) == []
            and holds_whole("p-ada", TIDE) == []
            and holds_whole("p-ada", PANTRY) == [],
            (holds("p-ada", CODE_TWO), reply3["content"][:100]),
        )
        refused = (
            ada.patch(f"/api/documents/folders/{grant}", json={"includeHidden": True}).status_code,
            ada.post(f"/api/documents/folders/{grant}/sync", json={}).status_code,
            ada.post("/api/documents/folders", json={"grantId": grant}).status_code,
        )
        check(
            "it cannot be changed, synced or added again: the grant is gone",
            refused == (409, 409, 404),
            refused,
        )
        check(
            "the folder's files are exactly as they were when it was withdrawn",
            snapshot(folder) == folder_before,
        )
        check(
            "she dismisses it, and it is gone from her list",
            ada.delete(f"/api/documents/folders/{grant}").status_code == 204
            and ada.folders()["folders"] == []
            and ada.get(f"/api/documents/folders/{grant}").status_code == 404,
        )

    except Exception as exc:
        traceback.print_exc()
        check(f"the run went to plan ({type(exc).__name__}: {exc})", False)
    finally:
        say("11. teardown by pid")
        if server is not None:
            server.should_exit = True
        for proc in procs:
            subprocess.run(
                ["taskkill", "/PID", str(proc.pid), "/T", "/F"], capture_output=True, check=False
            )
        time.sleep(2)
        # The link is removed first: removing a folder that holds one must not reach what it points at.
        with contextlib.suppress(OSError):
            (folder / "shortcut").rmdir()
        still = [p for p in OWNED if listening(p) and p not in OURS]
        check("no owned process port is still listening", not still, still)
        print(f"\nlogs: {logs}\n{PASSES} passed, {FAILURES} failed", flush=True)
    return 1 if FAILURES else 0


if __name__ == "__main__":
    raise SystemExit(main())
