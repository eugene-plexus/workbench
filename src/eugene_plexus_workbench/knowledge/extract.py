"""Reading a file: the text in it, and the truth when there is none (MR2).

This opens a person's file, and a file is someone else's bytes: a PDF can be
hostile or merely broken, a Word file can expand a thousandfold. So it never
runs in Workbench's own process. `worker.py` runs it in a short-lived child
(`reader.py` starts that child, times it and ends it), and what comes back is
the text with where it came from -- a page at a time for a PDF, with the
headings Markdown, Word and HTML say -- or the reason there is no text, in the
words a person reads: "this PDF has no text, only pictures of pages", "this file
is password-protected", and, when it is damaged, the cause as the library
observed it.

Nothing here decides what a model is given of a file (`context.py`); it says
what the file says. The text is made safe to carry, because it goes into a
model's prompt: control characters and invisible formatting characters (the
bidirectional overrides and tag characters used to hide words from a person)
are removed and counted, a lone surrogate (which no JSON can carry) becomes
the replacement character, and HTML's scripts, styles, comments and hidden
elements are dropped.

The libraries are imported where they are used, so the app can import this
module for its constants and its rendering without loading a PDF reader.

**PDF**: pypdfium2 (BSD-3 and Apache-2.0), with pypdf (BSD-3) as the fallback
for what PDFium refuses. Never PyMuPDF: it is AGPL, and Eugene is Apache-2.0.
**Word**: python-docx (MIT). **Everything else**: the standard library.
"""

from __future__ import annotations

import base64
import codecs
import contextlib
import html
import re
import struct
import zipfile
import zlib
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

#: The most text read from one file, in characters (about a million tokens, more
#: than any model's window). A file past it is read this far, and says so.
MAX_TEXT_CHARS = 3_000_000
#: Of a text file, the bytes read: UTF-8 needs up to four for a character.
MAX_READ_BYTES = 12_000_000
MAX_PAGES = 5_000
#: A Word file is a ZIP; what it holds, once opened, is bounded before it is.
MAX_ZIP_ENTRIES = 20_000
MAX_UNCOMPRESSED = 300 * 1024 * 1024
MAX_DOCUMENT_XML = 100 * 1024 * 1024
MAX_HEADINGS = 2_000
#: A PDF with fewer characters than this per page, on average, has no text layer.
MIN_CHARS_PER_PAGE = 5

KINDS = ("pdf", "docx", "markdown", "csv", "html", "code", "text", "image")
LABELS = {
    "pdf": "PDF",
    "docx": "Word",
    "markdown": "Markdown",
    "csv": "CSV",
    "html": "HTML",
    "code": "code",
    "text": "text",
    "image": "picture",
}
#: Kinds whose text is the file's own, so a position in it is a line of the file.
LINE_KINDS = ("markdown", "csv", "code", "text")


# --------------------------------------------------------------------------- #
# what a read comes to
# --------------------------------------------------------------------------- #


@dataclass(frozen=True)
class Part:
    """One stretch of a file in reading order: a PDF's page (`page` from 1),
    or the whole text of any other kind (`page` None)."""

    text: str
    page: int | None = None


@dataclass(frozen=True)
class Heading:
    level: int
    text: str
    #: Line of the single part's text it is on, from 1.
    line: int


@dataclass
class Extracted:
    kind: str
    parts: list[Part]
    language: str | None = None
    #: A PDF's real page count, including pages that had no text.
    pages: int | None = None
    #: A line-based kind's line count (`LINE_KINDS`).
    lines: int | None = None
    headings: list[Heading] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    #: The reader stopped at its own cap (`MAX_TEXT_CHARS`, `MAX_PAGES`).
    truncated: bool = False
    read_by: str | None = None

    def to_json(self) -> dict[str, Any]:
        return {
            "kind": self.kind,
            "language": self.language,
            "pages": self.pages,
            "lines": self.lines,
            "headings": [[h.level, h.text, h.line] for h in self.headings],
            "notes": self.notes,
            "truncated": self.truncated,
            "readBy": self.read_by,
            "parts": [[p.page, p.text] for p in self.parts],
        }

    @classmethod
    def from_json(cls, data: Any) -> Extracted:
        """What `to_json` wrote, or `ValueError` for anything else."""
        try:
            kind = data["kind"]
            if kind not in KINDS:
                raise ValueError(f"unknown kind {kind!r}")
            parts = [
                Part(text=str(t), page=None if p is None else int(p)) for p, t in data["parts"]
            ]
            return cls(
                kind=kind,
                parts=parts,
                language=data["language"],
                pages=data["pages"],
                lines=data["lines"],
                headings=[Heading(int(a), str(b), int(c)) for a, b, c in data["headings"]],
                notes=[str(n) for n in data["notes"]],
                truncated=bool(data["truncated"]),
                read_by=data.get("readBy"),
            )
        except (KeyError, TypeError, ValueError) as exc:
            raise ValueError(f"not a saved reading: {exc}") from exc


class Unreadable(Exception):
    """A file that gives no text, and why: `code` for a program, `reason` for a
    person. `state` is `skipped` (nothing to read, or nothing allowed) or
    `failed` (it could not be read)."""

    def __init__(
        self,
        code: str,
        reason: str,
        *,
        state: str = "failed",
        pages: int | None = None,
        detail: str | None = None,
    ) -> None:
        super().__init__(reason)
        self.code = code
        self.reason = reason
        self.state = state
        self.pages = pages
        self.detail = detail

    @property
    def message(self) -> str:
        return f"{'Skipped' if self.state == 'skipped' else 'Failed'}: {self.reason}"

    def to_json(self) -> dict[str, Any]:
        return {
            "code": self.code,
            "reason": self.reason,
            "state": self.state,
            "pages": self.pages,
            "detail": self.detail,
        }


def say(text: object, limit: int = 240) -> str:
    """An observed cause as one short line: whitespace collapsed, cut with a mark."""
    one = " ".join(str(text).split())
    return one if len(one) <= limit else one[: limit - 1].rstrip() + "…"


def render_page(path: Path, number: int, words: str = "") -> dict[str, Any]:
    """PDFium runs only in the restricted reader child. Return a bounded PNG and boxes."""
    import pypdfium2 as pdfium

    from .. import files

    with path.open("rb") as source:
        head = source.read(files.HEAD_BYTES)
    detected = files.detect(head, None, path.name)
    if detected.kind == "image":
        if number != 1:
            raise Unreadable("page_missing", "This picture has only one page.")
        if path.stat().st_size > files.LIMITS["image"]:
            raise Unreadable(
                "image_too_large", "This picture exceeds the gateway's 5 MiB image limit."
            )
        return {
            "image": f"data:{detected.media_type};base64,"
            + base64.b64encode(path.read_bytes()).decode("ascii"),
            "pageCount": 1,
            "width": 0,
            "height": 0,
            "boxes": [],
        }

    with pdfium.PdfDocument(path) as pdf:
        if not 1 <= number <= len(pdf):
            raise Unreadable("page_missing", f"This PDF has {len(pdf)} pages, not page {number}.")
        with contextlib.closing(pdf[number - 1]) as page:
            pw, ph = page.get_size()
            scale = min(150 / 72, 2000 / max(pw, ph))
            with contextlib.closing(page.render(scale=scale, rev_byteorder=True)) as bitmap:
                array = bitmap.to_numpy()
                height, width = array.shape[:2]
                # PDFium's RGB output; encode PNG using the standard library.
                rgb = array[:, :, :3].copy()
                rows = b"".join(b"\0" + row.tobytes() for row in rgb)

                def block(kind: bytes, body: bytes) -> bytes:
                    return (
                        struct.pack(">I", len(body))
                        + kind
                        + body
                        + struct.pack(">I", zlib.crc32(kind + body))
                    )

                png = (
                    b"\x89PNG\r\n\x1a\n"
                    + block(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
                    + block(b"IDAT", zlib.compress(rows))
                    + block(b"IEND", b"")
                )
            boxes: list[list[float]] = []
            if words.strip():
                with contextlib.closing(page.get_textpage()) as textpage:
                    # Search each normalised line as PDFium represents it. A chunk may
                    # cross pages; only words actually on this page receive boxes.
                    for line in words.splitlines():
                        phrase = " ".join(line.split())
                        if not phrase:
                            continue
                        search = textpage.search(phrase)
                        try:
                            while match := search.get_next():
                                start, count = match
                                for i in range(textpage.count_rects(start, count)):
                                    left, bottom, right, top = textpage.get_rect(i)
                                    boxes.append(
                                        [
                                            left * scale,
                                            (ph - top) * scale,
                                            (right - left) * scale,
                                            (top - bottom) * scale,
                                        ]
                                    )
                        finally:
                            search.close()
            return {
                "image": "data:image/png;base64," + base64.b64encode(png).decode("ascii"),
                "width": width,
                "height": height,
                "boxes": boxes,
                "pageCount": len(pdf),
            }


# --------------------------------------------------------------------------- #
# the text, as a model is shown it
# --------------------------------------------------------------------------- #

NO_TEXT_ON_PAGE = "(no text on this page)"


def page_marker(number: int) -> str:
    return f"[page {number}]"


def render_part(part: Part) -> str:
    """One part as the model reads it: a PDF's page opens with its marker, so
    a model can say which page a fact is on."""
    if part.page is None:
        return part.text
    body = part.text.strip()
    marker = page_marker(part.page)
    return f"{marker}\n{body}" if body else f"{marker} {NO_TEXT_ON_PAGE}"


def render(parts: list[Part]) -> str:
    """Every part, in order, as a model reads them whole."""
    return "\n\n".join(render_part(p) for p in parts)


# --------------------------------------------------------------------------- #
# making text safe to carry
# --------------------------------------------------------------------------- #

#: Controls (not tab, line feed or form feed), DEL and the C1 controls; and the
#: invisible formatting characters: soft hyphen, Arabic marks, the Mongolian
#: vowel separator, zero-width space, word joiner and invisible operators, the
#: bidirectional embeddings, overrides and isolates, the byte-order mark and
#: the noncharacters, interlinear annotation, and the tag characters. The
#: joiners and marks a script needs (U+200C to U+200F) stay.
_INVISIBLE = re.compile(
    "[\x00-\x08\x0b\x0e-\x1f\x7f-\x9f"
    "\U000000ad\U00000600-\U00000605\U0000061c\U000006dd\U0000070f\U000008e2\U0000180e"
    "\U0000200b\U00002060-\U00002064\U00002066-\U0000206f"
    "\U0000202a-\U0000202e\U0000feff\U0000fff9-\U0000fffb\U0000fffe\U0000ffff"
    "\U000110bd\U000110cd\U0001bca0-\U0001bca3\U0001d173-\U0001d17a\U000e0001\U000e0020-\U000e007f]"
)
_SURROGATE = re.compile("[\ud800-\udfff]")
_BLANKS = re.compile(r"\n{3,}")
_TRAILING = re.compile(r"[ \t]+\n")


def clean(text: str, *, layout: bool) -> tuple[str, int]:
    """`text` fit to put in a prompt, and how many characters were taken out.

    Line endings become `\\n`. With `layout` (text a PDF, Word or HTML reader
    made from a page, so its spacing is the layout's and not the author's),
    trailing spaces go and a run of blank lines is one."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\x0c", "\n").replace("\U00002028", "\n").replace("\U00002029", "\n")
    text, removed = _INVISIBLE.subn("", text)
    text = _SURROGATE.sub("\U0000fffd", text)
    if layout:
        text = _TRAILING.sub("\n", text)
        text = _BLANKS.sub("\n\n", text).strip("\n")
    return text, removed


def _removed_note(count: int) -> list[str]:
    if not count:
        return []
    return [f"Removed {count:,} control or invisible character{'s' if count != 1 else ''}."]


# --------------------------------------------------------------------------- #
# decoding text
# --------------------------------------------------------------------------- #


def decode(data: bytes, *, complete: bool = True) -> tuple[str, list[str]]:
    """Text from bytes: UTF-8 (with or without a byte-order mark), UTF-16 with
    one, else Windows-1252, which says so. `complete` is False for a prefix of
    a file, whose last character may be cut in two."""
    notes: list[str] = []
    if data.startswith(codecs.BOM_UTF8):
        text = data[len(codecs.BOM_UTF8) :].decode("utf-8", errors="replace")
        return text, notes
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        if len(data) % 2 and not complete:
            data = data[:-1]
        return data.decode("utf-16", errors="replace"), notes
    decoder = codecs.getincrementaldecoder("utf-8")()
    try:
        return decoder.decode(data, final=complete), notes
    except UnicodeDecodeError:
        pass
    notes.append("This file is not UTF-8, so Workbench read it as Windows-1252 text.")
    return data.decode("cp1252", errors="replace"), notes


# --------------------------------------------------------------------------- #
# headings
# --------------------------------------------------------------------------- #

_ATX = re.compile(r"^ {0,3}(#{1,6})[ \t]+(.*?)(?:[ \t]+#+)?[ \t]*$")
_SETEXT = re.compile(r"^ {0,3}(=+|-+)[ \t]*$")


def markdown_headings(text: str) -> list[Heading]:
    """The headings of Markdown (ATX `#`, and setext underlines), outside code fences."""
    found: list[Heading] = []
    fence: str | None = None
    previous = ""
    for number, line in enumerate(text.split("\n"), start=1):
        stripped = line.lstrip()
        if stripped.startswith(("```", "~~~")):
            marker = stripped[:3]
            if fence is None:
                fence = marker
            elif marker == fence:
                fence = None
            previous = ""
            continue
        if fence is not None:
            continue
        atx = _ATX.match(line)
        if atx is not None and atx.group(2).strip():
            found.append(Heading(len(atx.group(1)), atx.group(2).strip(), number))
        else:
            under = _SETEXT.match(line)
            if (
                under is not None
                and len(line.strip()) >= 3
                and previous.strip()
                and not previous.lstrip().startswith((">", "-", "*", "+", "|", "#"))
            ):
                found.append(
                    Heading(1 if line.strip()[0] == "=" else 2, previous.strip(), number - 1)
                )
        previous = line
        if len(found) >= MAX_HEADINGS:
            break
    return found


# --------------------------------------------------------------------------- #
# plain text, Markdown, CSV, code
# --------------------------------------------------------------------------- #


def line_count(text: str) -> int:
    """How many lines `text` has: a last line without a line feed counts, and a
    line feed at the very end does not start another."""
    return text.count("\n") + (0 if text.endswith("\n") or not text else 1)


def read_text(path: Path, kind: str, language: str | None) -> Extracted:
    size = path.stat().st_size
    with path.open("rb") as handle:
        data = handle.read(MAX_READ_BYTES)
    truncated = len(data) < size
    text, notes = decode(data, complete=not truncated)
    text, removed = clean(text, layout=False)
    if len(text) > MAX_TEXT_CHARS:
        text, truncated = text[:MAX_TEXT_CHARS], True
    if not text.strip():
        raise Unreadable("empty", "this file has no text in it.", state="skipped")
    if truncated:
        notes.append(
            f"Workbench read the first {len(text):,} characters of this file; "
            "the rest was not read."
        )
    notes += _removed_note(removed)
    return Extracted(
        kind=kind,
        parts=[Part(text)],
        language=language,
        lines=line_count(text),
        headings=markdown_headings(text) if kind == "markdown" else [],
        notes=notes,
        truncated=truncated,
    )


# --------------------------------------------------------------------------- #
# HTML
# --------------------------------------------------------------------------- #

_VOID = frozenset(
    {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param"}
    | {"source", "track", "wbr"}
)
#: Whose contents are never text a person reads.
_SKIPPED = frozenset(
    {"script", "style", "noscript", "template", "svg", "canvas", "iframe", "object"}
    | {"embed", "audio", "video", "map"}
)
_BLOCKS = frozenset(
    {"address", "article", "aside", "blockquote", "details", "dd", "div", "dl", "dt"}
    | {"fieldset", "figcaption", "figure", "footer", "form", "header", "main", "nav", "ol"}
    | {"p", "section", "summary", "ul"}
)
#: Of those, the ones a blank line sets apart; the rest only end the line.
_PARAGRAPHS = frozenset({"p", "blockquote", "section", "article"})
_HIDDEN_STYLE = re.compile(r"(?:display\s*:\s*none|visibility\s*:\s*hidden)", re.IGNORECASE)
_CHARSET = re.compile(rb"charset\s*=\s*[\"']?\s*([A-Za-z0-9_.:-]+)", re.IGNORECASE)


def _hidden(attrs: list[tuple[str, str | None]]) -> bool:
    for name, value in attrs:
        if name == "hidden" or (name == "aria-hidden" and (value or "").lower() == "true"):
            return True
        if name == "style" and value and _HIDDEN_STYLE.search(value):
            return True
        if name == "type" and (value or "").lower() == "hidden":
            return True
    return False


class _HtmlText(HTMLParser):
    """HTML as the text a person reads: headings as `#`, lists as `-`, tables as
    rows; scripts, styles, comments and hidden elements dropped."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.title = ""
        self._in_title = False
        self._in_head = False
        self._pre = 0
        self._skip: str | None = None
        self._depth = 0
        self._lists: list[list[int]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None
        self._rows_done = 0
        #: Trailing line feeds in the output, and whether it ends in a space or
        #: a line feed (or is empty), so that spacing is decided once.
        self._newlines = 0
        self._spaced = True

    # --- output ---------------------------------------------------------

    def _emit(self, text: str) -> None:
        if not text:
            return
        self.parts.append(text)
        body = text.rstrip("\n")
        self._newlines = len(text) - len(body) if body else self._newlines + len(text)
        self._spaced = text[-1] in " \n"

    def _newline(self, count: int) -> None:
        if self.parts and self._newlines < count:
            self._emit("\n" * (count - self._newlines))

    def _write(self, text: str) -> None:
        if self._cell is not None:
            self._cell.append(text)
        else:
            self._emit(text)

    def _flush_row(self) -> None:
        if self._row:
            self._newline(1)
            self._emit("| " + " | ".join(self._row) + " |\n")
            if self._rows_done == 0:
                self._emit("| " + " | ".join("---" for _ in self._row) + " |\n")
            self._rows_done += 1
        self._row = None

    # --- tags -----------------------------------------------------------

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if self._skip is not None:
            if tag == self._skip and tag not in _VOID:
                self._depth += 1
            return
        if tag in _SKIPPED or (_hidden(attrs) and tag not in _VOID):
            self._skip, self._depth = tag, 1
            return
        if tag == "head":
            self._in_head = True
        elif tag == "title" and self._in_head:
            self._in_title = True
        elif tag == "pre":
            self._newline(2)
            self._pre += 1
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self._newline(2)
            self._write("#" * int(tag[1]) + " ")
        elif tag in ("ul", "ol"):
            self._newline(2 if not self._lists else 1)
            self._lists.append([0] if tag == "ol" else [-1])
        elif tag == "li":
            self._newline(1)
            counter = self._lists[-1] if self._lists else [-1]
            indent = "  " * max(0, len(self._lists) - 1)
            if counter[0] >= 0:
                counter[0] += 1
                self._write(f"{indent}{counter[0]}. ")
            else:
                self._write(f"{indent}- ")
        elif tag == "br":
            self._write(" " if self._cell is not None else "\n")
        elif tag == "hr":
            self._newline(2)
            self._write("---")
            self._newline(2)
        elif tag == "table":
            self._flush_row()
            self._newline(2)
            self._rows_done = 0
        elif tag == "tr":
            self._flush_row()
            self._row = []
        elif tag in ("td", "th"):
            if self._row is None:
                self._row = []
            self._cell = []
        elif tag in _BLOCKS:
            self._newline(2 if tag in _PARAGRAPHS else 1)

    def handle_endtag(self, tag: str) -> None:
        if self._skip is not None:
            if tag == self._skip:
                self._depth -= 1
                if self._depth == 0:
                    self._skip = None
            return
        if tag == "head":
            self._in_head = False
        elif tag == "title":
            self._in_title = False
        elif tag == "pre":
            self._pre = max(0, self._pre - 1)
            self._newline(2)
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self._newline(2)
        elif tag in ("ul", "ol"):
            if self._lists:
                self._lists.pop()
            self._newline(2 if not self._lists else 1)
        elif tag in ("td", "th"):
            if self._cell is not None and self._row is not None:
                self._row.append(" ".join("".join(self._cell).split()).replace("|", "\\|"))
            self._cell = None
        elif tag == "tr":
            self._flush_row()
        elif tag == "table":
            self._flush_row()
            self._newline(2)
        elif tag in _BLOCKS:
            self._newline(2 if tag in _PARAGRAPHS else 1)

    def handle_data(self, data: str) -> None:
        if self._skip is not None:
            return
        if self._in_title:
            self.title += data
            return
        if self._in_head:
            return
        if self._pre or self._cell is not None:
            self._write(data)
            return
        words = " ".join(data.split())
        if not words:
            if data and not self._spaced:
                self._emit(" ")
            return
        if data[0].isspace() and not self._spaced:
            words = " " + words
        if data[-1].isspace():
            words += " "
        self._emit(words)

    def handle_comment(self, data: str) -> None:  # a comment is never text a person reads
        return

    def text(self) -> str:
        self._flush_row()
        return "".join(self.parts)


def _declared_charset(data: bytes) -> str | None:
    found = _CHARSET.search(data[:4096])
    if found is None:
        return None
    name = found.group(1).decode("ascii", errors="ignore")
    try:
        return codecs.lookup(name).name
    except LookupError:
        return None


def read_html(path: Path) -> Extracted:
    size = path.stat().st_size
    with path.open("rb") as handle:
        data = handle.read(MAX_READ_BYTES)
    truncated = len(data) < size
    text, notes = decode(data, complete=not truncated)
    if notes:
        declared = _declared_charset(data)
        if declared is not None:
            with contextlib.suppress(LookupError):
                text, notes = data.decode(declared, errors="replace"), []
    parser = _HtmlText()
    parser.feed(text)
    parser.close()
    body = parser.text()
    title = " ".join(parser.title.split())
    if title and not re.match(rf"\s*#\s+{re.escape(title)}\s*$", body.split("\n", 1)[0]):
        body = f"# {title}\n\n{body}"
    body, removed = clean(body, layout=True)
    if len(body) > MAX_TEXT_CHARS:
        body, truncated = body[:MAX_TEXT_CHARS], True
    if not body.strip():
        raise Unreadable("empty", "this web page has no text in it.", state="skipped")
    if truncated:
        notes.append(
            f"Workbench read the first {len(body):,} characters of this file; "
            "the rest was not read."
        )
    notes += _removed_note(removed)
    return Extracted(
        kind="html",
        parts=[Part(body)],
        headings=markdown_headings(body),
        notes=notes,
        truncated=truncated,
    )


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #

#: PDFium's own code for a file it cannot open without a password.
_PDFIUM_PASSWORD = 4
_PDFIUM_LOAD = re.compile(r"^Failed to load document \((.*)\)\.?$")


def _pdfium_words(exc: BaseException) -> str:
    """What PDFium said, once: `PDFium: Data format error`."""
    said = say(exc)
    found = _PDFIUM_LOAD.match(said)
    return found.group(1) if found else f"PDFium: {said}"


def _ranges(numbers: list[int]) -> str:
    """Page numbers as a person says them: `2, 4 to 9 and 12`."""
    spans: list[str] = []
    start = last = numbers[0]
    for n in numbers[1:]:
        if n == last + 1:
            last = n
            continue
        spans.append(str(start) if start == last else f"{start} to {last}")
        start = last = n
    spans.append(str(start) if start == last else f"{start} to {last}")
    return spans[0] if len(spans) == 1 else ", ".join(spans[:-1]) + " and " + spans[-1]


def _pdfium_pages(path: Path) -> tuple[list[str], int, bool]:
    """(each page's text, the page count, whether the read stopped early) by
    PDFium. Raises `Unreadable` for a password, `_Refused` for a file PDFium
    will not open (the caller tries pypdf)."""
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(str(path))
    except pdfium.PdfiumError as exc:
        if getattr(exc, "err_code", None) == _PDFIUM_PASSWORD:
            raise Unreadable(
                "password", "this file is password-protected.", state="skipped"
            ) from exc
        raise _Refused(_pdfium_words(exc)) from exc
    texts: list[str] = []
    chars = 0
    stopped = False
    try:
        count = len(pdf)
        for index in range(count):
            if index >= MAX_PAGES or chars > MAX_TEXT_CHARS:
                stopped = True
                break
            page = pdf[index]
            try:
                textpage = page.get_textpage()
                try:
                    text = textpage.get_text_bounded()
                finally:
                    textpage.close()
            finally:
                page.close()
            texts.append(text)
            chars += len(text)
    finally:
        pdf.close()
    return texts, count, stopped


class _Refused(Exception):
    """PDFium would not open the file; `args[0]` is what it said."""


def _pypdf_pages(path: Path, first: str) -> tuple[list[str], int, bool]:
    """The same by pypdf, for a file PDFium refused (`first` is what PDFium said)."""
    import pypdf

    try:
        reader = pypdf.PdfReader(str(path), strict=False)
        if reader.is_encrypted and not reader.decrypt(""):
            raise Unreadable("password", "this file is password-protected.", state="skipped")
        count = len(reader.pages)
    except Unreadable:
        raise
    except Exception as exc:
        raise Unreadable(
            "damaged",
            f"Workbench could not read this file ({first}; pypdf: {say(f'{exc}')}).",
            detail=f"{type(exc).__name__}: {exc}",
        ) from exc
    texts: list[str] = []
    chars = 0
    stopped = False
    for index in range(count):
        if index >= MAX_PAGES or chars > MAX_TEXT_CHARS:
            stopped = True
            break
        try:
            text = reader.pages[index].extract_text() or ""
        except Exception:
            text = ""
        texts.append(text)
        chars += len(text)
    return texts, count, stopped


def read_pdf(path: Path) -> Extracted:
    with path.open("rb") as handle:
        head = handle.read(1024)
    if b"%PDF-" not in head:
        raise Unreadable(
            "damaged",
            "Workbench could not read this file (it does not start like a PDF).",
            detail="no %PDF header",
        )
    try:
        texts, count, stopped = _pdfium_pages(path)
    except _Refused as refused:
        texts, count, stopped = _pypdf_pages(path, str(refused))
    cleaned: list[str] = []
    removed = 0
    for text in texts:
        page, gone = clean(text, layout=True)
        cleaned.append(page)
        removed += gone
    if count == 0:
        raise Unreadable("empty", "this PDF has no pages in it.", state="skipped", pages=0)
    letters = sum(len("".join(t.split())) for t in cleaned)
    pages_read = len(cleaned)
    if letters == 0 or letters < MIN_CHARS_PER_PAGE * pages_read:
        raise Unreadable(
            "no_text",
            "this PDF has no text, only pictures of pages.",
            state="skipped",
            pages=count,
        )
    notes: list[str] = []
    empty = [n for n, t in enumerate(cleaned, start=1) if not t.strip()]
    if len(empty) == 1:
        notes.append(f"Page {empty[0]} has no text (it may be a picture).")
    elif empty:
        notes.append(f"Pages {_ranges(empty)} have no text (they may be pictures).")
    if stopped:
        notes.append(
            f"Workbench read the first {pages_read:,} of {count:,} pages; the rest were not read."
        )
    notes += _removed_note(removed)
    return Extracted(
        kind="pdf",
        parts=[Part(text, page=n) for n, text in enumerate(cleaned, start=1)],
        pages=count,
        notes=notes,
        truncated=stopped,
    )


# --------------------------------------------------------------------------- #
# Word
# --------------------------------------------------------------------------- #

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_HEADING = re.compile(r"^heading\s*(\d)$", re.IGNORECASE)


def _check_package(path: Path) -> None:
    """Refuse a Word file by what it declares, before anything is expanded: how
    many parts, how large they would be, and that a document is in it."""
    try:
        with zipfile.ZipFile(path) as package:
            infos = package.infolist()
    except zipfile.BadZipFile as exc:
        raise Unreadable(
            "damaged",
            f"Workbench could not read this file (it is not a readable Word file: {say(exc)}).",
            detail=f"BadZipFile: {exc}",
        ) from exc
    if len(infos) > MAX_ZIP_ENTRIES:
        raise Unreadable(
            "too_big",
            f"this Word file holds {len(infos):,} parts, more than Workbench opens "
            f"({MAX_ZIP_ENTRIES:,}).",
            state="skipped",
        )
    total = sum(i.file_size for i in infos)
    if total > MAX_UNCOMPRESSED:
        raise Unreadable(
            "too_big",
            f"this Word file expands to {total // (1024 * 1024):,} MB when opened, more than "
            f"Workbench opens ({MAX_UNCOMPRESSED // (1024 * 1024)} MB).",
            state="skipped",
        )
    document = next((i for i in infos if i.filename == "word/document.xml"), None)
    if document is None:
        raise Unreadable(
            "damaged",
            "Workbench could not read this file (it is a ZIP file with no Word document in it).",
            detail="no word/document.xml",
        )
    if document.file_size > MAX_DOCUMENT_XML:
        raise Unreadable(
            "too_big",
            f"this Word file's text is {document.file_size // (1024 * 1024):,} MB when opened, "
            f"more than Workbench opens ({MAX_DOCUMENT_XML // (1024 * 1024)} MB).",
            state="skipped",
        )


def _heading_level(paragraph: Any) -> int | None:
    try:
        style = paragraph.style
        name = (style.name or "") if style is not None else ""
        style_id = (style.style_id or "") if style is not None else ""
    except Exception:
        return None
    if name.lower() == "title":
        return 1
    found = _HEADING.match(name) or re.match(r"^Heading(\d)$", style_id)
    return min(6, max(1, int(found.group(1)))) if found else None


def _list_depth(paragraph: Any) -> int | None:
    """How deep a list item is (0 for the first level), or None when the
    paragraph is not one: a list style, or direct list formatting."""
    properties = paragraph._p.pPr  # the paragraph's own formatting, or None
    numbering = properties.numPr if properties is not None else None
    if numbering is not None:
        level = numbering.ilvl.val if numbering.ilvl is not None else 0
        return int(level or 0)
    try:
        name = (paragraph.style.name or "") if paragraph.style is not None else ""
    except Exception:
        return None
    found = re.match(r"^List (?:Bullet|Number|Paragraph|Continue)(?: (\d))?$", name)
    return int(found.group(1) or 1) - 1 if found else None


def _numbered(paragraph: Any) -> bool:
    try:
        name = (paragraph.style.name or "") if paragraph.style is not None else ""
    except Exception:
        return False
    return name.startswith("List Number")


def _table_text(table: Any) -> list[str]:
    rows: list[list[str]] = []
    for row in table.rows:
        cells: list[str] = []
        seen: set[int] = set()
        for cell in row.cells:
            if id(cell._tc) in seen:  # a merged cell shows once
                continue
            seen.add(id(cell._tc))
            words = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            cells.append(words.replace("|", "\\|"))
        rows.append(cells)
    if not rows:
        return []
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join("---" for _ in rows[0]) + " |"]
    lines += ["| " + " | ".join(row) + " |" for row in rows[1:]]
    return lines


def read_docx(path: Path) -> Extracted:
    _check_package(path)
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise Unreadable(
            "damaged",
            f"Workbench could not read this file ({say(f'{type(exc).__name__}: {exc}')}).",
            detail=f"{type(exc).__name__}: {exc}",
        ) from exc

    def blocks(parent: Any) -> Any:
        for child in parent.iterchildren():
            if child.tag == _W + "p":
                yield Paragraph(child, document)
            elif child.tag == _W + "tbl":
                yield Table(child, document)
            elif child.tag == _W + "sdt":
                content = child.find(_W + "sdtContent")
                if content is not None:
                    yield from blocks(content)

    lines: list[str] = []
    counters: dict[int, int] = {}
    in_list = False

    def block(*text: str, tight: bool = False) -> None:
        if lines and not tight and lines[-1] != "":
            lines.append("")
        lines.extend(text)

    for item in blocks(document.element.body):
        if isinstance(item, Table):
            counters.clear()
            in_list = False
            rows = _table_text(item)
            if rows:
                block(*rows)
            continue
        text = item.text.strip()
        level = _heading_level(item)
        if level is not None and text:
            counters.clear()
            in_list = False
            block("#" * level + " " + text)
            continue
        depth = _list_depth(item)
        if depth is not None and text:
            indent = "  " * depth
            if _numbered(item):
                counters[depth] = counters.get(depth, 0) + 1
                marker = f"{counters[depth]}."
            else:
                counters.pop(depth, None)
                marker = "-"
            block(f"{indent}{marker} {text}", tight=in_list)
            in_list = True
            continue
        counters.clear()
        if text:
            in_list = False
            block(text)
    text, removed = clean("\n".join(lines), layout=True)
    truncated = len(text) > MAX_TEXT_CHARS
    if truncated:
        text = text[:MAX_TEXT_CHARS]
    if not text.strip():
        raise Unreadable("empty", "this Word file has no text in it.", state="skipped")
    notes = _removed_note(removed)
    if truncated:
        notes.append(
            f"Workbench read the first {len(text):,} characters of this file; "
            "the rest was not read."
        )
    return Extracted(
        kind="docx",
        parts=[Part(text)],
        headings=markdown_headings(text),
        notes=notes,
        truncated=truncated,
    )


# --------------------------------------------------------------------------- #
# the entry
# --------------------------------------------------------------------------- #


def read(path: Path, kind: str, language: str | None = None) -> Extracted:
    """Read `path` as a file of this kind (`KINDS`). Raises `Unreadable`."""
    if kind == "pdf":
        return read_pdf(path)
    if kind == "docx":
        return read_docx(path)
    if kind == "html":
        return read_html(path)
    if kind in ("markdown", "csv", "code", "text"):
        return read_text(path, kind, language)
    raise Unreadable("unsupported", f"Workbench does not read {html.escape(kind)} files.")
